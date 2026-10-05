"""Sample business documents for the demo shop (Gupta Digital World), one per
supported format. The operations log uses the dates of the events planted in
the demo sales, so "why was that day so slow?" has an answer in the documents.

    python -m legacy.rag.samples data/samples/docs   # write them to a folder
"""

from __future__ import annotations

import io
import sys
from datetime import date, timedelta
from pathlib import Path

from legacy.datagen import planted_events

RETURNS_POLICY = """# Returns, refunds and warranty policy

Gupta Digital World - applies to all stores, the website and marketplace orders.

## Returns

Customers can return a product within 7 days of purchase with the original bill or order number.
Accessories (cases, screen guards, cables, chargers) must be unused and in the original packaging.
Opened electronics - phones, laptops, tablets, watches, earbuds - can be returned only if they are faulty
on arrival; otherwise we offer an exchange within 7 days.
Marketplace orders are returned through the marketplace's own process.

## Refunds

Refunds go back to the original payment method: UPI and card refunds take 5 to 7 working days,
cash purchases are refunded in cash at the store the same day.
Discounts given on the original bill are deducted from the refund.

## Warranty

Phones, laptops and tablets carry the manufacturer's 1 year warranty; we register the product
for the customer at billing. Accessories carry a 6 month replacement warranty at the store.
Extended warranty (1 extra year) can be bought within 30 days of purchase for 8% of the product price.

## Who decides

Store managers can approve returns outside these rules up to Rs 5,000. Anything larger needs the
owner's approval (Mr. R. Gupta).
"""

PRICING_POLICY = [
    ("Title", "Discount and pricing policy"),
    ("Heading 1", "Everyday discounts"),
    (None, "Staff may give up to 5% off any bill without approval. Discounts from 5% to 10% need the store "
           "manager's approval. Discounts above 10% are only allowed during announced sales."),
    (None, "Loyalty customers - more than 5 orders in the last 12 months - get an extra 5% on accessories."),
    ("Heading 1", "Festive season sale"),
    (None, "Every October and November we run the festive sale with 5%, 10% or 15% off selected items. "
           "The festive season is our busiest period: plan staff and stock for about 40% more orders "
           "than a normal month."),
    ("Heading 1", "Price changes"),
    (None, "Prices follow supplier costs. When a supplier raises prices we pass the increase on to "
           "customers from the date it takes effect; the last Electronics increase was 8%, following the "
           "SmartTech Distributors price revision (see the supplier terms)."),
    ("Heading 1", "Price list"),
    ("table", [("Product", "Price (Rs)", "Category"),
               ("Smartphone X1", "18,999", "Electronics"),
               ("Laptop 14 inch", "52,999", "Electronics"),
               ("Phone Case", "799", "Accessories"),
               ("Screen Guard", "349", "Accessories"),
               ("Air Fryer", "6,999", "Home Appliances")]),
]

SUPPLIER_TERMS = [
    "Supplier terms summary - SmartTech Distributors Pvt Ltd",
    "",
    "Supplies: smartphones, tablets, laptops, smart watches, earbuds and speakers (all Electronics).",
    "Payment terms: 30 days from invoice. 2% early-payment discount if paid within 10 days.",
    "Minimum order: Rs 2,00,000 per purchase order. Delivery lead time: 5 working days to any store.",
    "Defective units: returned to the supplier within 15 days of our purchase for a full credit.",
    "",
    "Price revision clause",
    "The supplier may revise prices once per quarter with 30 days notice. The latest revision raised",
    "all Electronics list prices by 8%, and our selling prices were raised by the same 8%.",
    "\f",
    "Contacts",
    "Account manager: Priya Nair, priya.nair@smarttech.example, +91 98200 00000.",
    "Escalations: operations desk, ops@smarttech.example.",
    "",
    "Other suppliers",
    "Accessories come from Kaveri Accessories (payment 15 days, lead time 3 days).",
    "Home appliances and personal care come from HomeLine Traders (payment 30 days).",
]


def sample_documents(end: date | None = None, months: int = 18) -> list[tuple[str, bytes]]:
    """(filename, content) pairs. ``end`` should be the demo data's last day."""
    events = planted_events(months, end)
    return [
        ("returns_and_warranty_policy.md", RETURNS_POLICY.encode()),
        ("discount_and_pricing_policy.docx", _docx(PRICING_POLICY)),
        ("supplier_terms_smarttech.pdf", _pdf(SUPPLIER_TERMS)),
        ("operations_log.txt", _operations_log(events).encode()),
    ]


def _operations_log(events: dict) -> str:
    flash, outage, rise = (events[k].date() for k in ("flash_sale", "outage", "price_rise"))
    fmt = lambda d: f"{d.day} {d:%B %Y}"  # noqa: E731
    return f"""Operations log - Gupta Digital World

{fmt(flash - timedelta(days=3))}: Marketing approved a one-day flash sale for {fmt(flash)}: 20% off \
everything, announced by SMS and on Instagram to all customers.

{fmt(flash)}: Flash sale day. Stores and the website were very busy - roughly three times the usual \
number of orders. Extra staff were called in at Mumbai and Pune.

{fmt(rise)}: Electronics prices raised by 8% in all channels after the SmartTech Distributors price \
revision. Accessories and appliances unchanged.

{fmt(outage)}: Billing system outage. The POS server failed in the morning and billing was down for \
most of the day in all stores; the website checkout was also affected. Only a small number of orders \
were billed. The IT vendor replaced the server's disk; systems were back the next morning.

{fmt(outage + timedelta(days=1))}: All systems working. Review meeting: a backup billing laptop will be \
kept at each store from now on.
"""


def _docx(content: list) -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.title = "Discount and pricing policy"
    for style, value in content:
        if style == "table":
            table = document.add_table(rows=0, cols=len(value[0]))
            for row in value:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
        elif style:
            document.add_paragraph(value, style=style)
        else:
            document.add_paragraph(value)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _pdf(lines: list[str]) -> bytes:
    """A minimal text PDF (Helvetica, one line per entry, "\\f" starts a new page)."""
    pages, page = [], []
    for line in lines:
        if line == "\f":
            pages.append(page)
            page = []
        else:
            page.append(line)
    pages.append(page)

    def escape(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects: list[bytes] = []
    n_pages = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    font_id = 3 + 2 * n_pages
    for i, text_lines in enumerate(pages):
        stream = "BT /F1 11 Tf 14 TL 56 780 Td " + " ".join(f"({escape(t)}) Tj T*" for t in text_lines) + " ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {4 + 2 * i} 0 R "
                       f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    out.write("".join(f"{o:010d} 00000 n \n" for o in offsets).encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


if __name__ == "__main__":
    folder = Path(sys.argv[1] if len(sys.argv) > 1 else "data/samples/docs")
    folder.mkdir(parents=True, exist_ok=True)
    for name, content in sample_documents():
        (folder / name).write_bytes(content)
        print(f"wrote {folder / name}")
