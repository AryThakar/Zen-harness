# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0
"""Small documents of every kind search reads (Markdown, Word, PDF, text), built in memory for the tests."""

from __future__ import annotations

import io

RETURNS = """# Returns, refunds and warranty policy

## Returns
Unopened products can be returned within 7 days of purchase with the bill.

## Refunds
Refunds go back to the original payment method within 5 working days.

## Warranty
Electronics carry the maker's warranty; we help with the claim.
"""


def pdf(lines: list[str]) -> bytes:
    """A minimal text PDF (Helvetica, one line per entry, "\\f" starts a new page)."""
    pages, page = [], []
    for line in lines:
        if line == "\f":
            pages.append(page)
            page = []
        else:
            page.append(line)
    pages.append(page)

    def escape(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font = 3 + 2 * len(pages)
    for i, text_lines in enumerate(pages):
        stream = "BT /F1 11 Tf 14 TL 56 780 Td " + " ".join(f"({escape(t)}) Tj T*" for t in text_lines) + " ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {4 + 2 * i} 0 R "
                       f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode())
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    out.write("".join(f"{o:010d} 00000 n \n" for o in offsets).encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def pricing_docx() -> bytes:
    import docx
    document = docx.Document()
    document.core_properties.title = "Discount and pricing policy"
    document.add_paragraph("Discounts", style="Heading 1")
    document.add_paragraph("Members get 5% off accessories all year.")
    document.add_paragraph("Price list", style="Heading 1")
    table = document.add_table(rows=0, cols=3)
    for row in (("Product", "Price", "Category"), ("Smartphone X1", "18,999", "Electronics")):
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


SUPPLIER_TERMS = ["Supplier terms - SmartTech Distributors", "Payment within 30 days of the invoice.", "\f",
                  "Account manager: Priya Nair", "Returns of faulty stock within 14 days."]
LOG = " ".join(f"Day {n}: the store opened on time and the billing system worked." for n in range(1, 15))


def samples() -> dict[str, bytes]:
    return {"returns_and_warranty_policy.md": RETURNS.encode(), "discount_and_pricing_policy.docx": pricing_docx(),
            "supplier_terms_smarttech.pdf": pdf(SUPPLIER_TERMS), "operations_log.txt": LOG.encode()}
