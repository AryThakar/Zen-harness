"""files() - main's one tool for the workspace's files: list, find, read, write, delete, download."""

from __future__ import annotations

import io

import pandas as pd
import pytest

from app.tools import approval
from app.tools.file_ops import PAGE, files, markdown_docx
from app.tools.files import find_files
from app.tools.search_index import SearchIndex
from app.tools.workspace import Workspace


@pytest.fixture
def ws(tmp_path):
    return Workspace(tmp_path / "ws")


@pytest.fixture
def index(tmp_path):
    return SearchIndex(tmp_path / "search.sqlite", embedder=None)  # keywords only


def run(ws, index, mode, **kw):
    return files(ws, index, mode, **kw)


def test_list_pages_newest_first(ws, index):
    assert run(ws, index, "list").result["items"] == [] and "empty" in run(ws, index, "list").result["note"]
    for n in range(PAGE + 2):
        ws.add_bytes("file", b"x", "txt", name=f"f{n}.txt")
    first = run(ws, index, "list").result
    assert first["total"] == PAGE + 2 and first["page"] == "1 of 2" and first["items"][0].startswith(f"f{PAGE + 1}.txt: ")
    assert first["next"] == "more: files(mode='list', page=2)"
    last = run(ws, index, "list", page=2).result
    assert len(last["items"]) == 2 and last["next"] == "end of list"
    assert run(ws, index, "list", page=3).status == "invalid_arguments"


def test_find_by_name_and_close_names_for_a_wrong_one(ws, index):
    ws.add_bytes("file", b"print(1)", "py", name="parabola.py")
    found = run(ws, index, "find", query="Parabola.PY").result
    assert found["found"] == 1 and found["files"][0]["path"] == "parabola.py" and found["files"][0]["match"].startswith("its name")
    near = run(ws, index, "find", query="parabla.py").result  # misspelt: no word of it is in a name
    assert near["found"] == 0 and "Did you mean: parabola.py?" in near["note"]
    assert run(ws, index, "find", name="parabola.py").status == "invalid_arguments"  # find takes query
    wrong = run(ws, index, "read", name="parabol.py")  # the closest names, never a listing line to copy back
    assert wrong.status == "invalid_arguments" and "Did you mean: parabola.py?" in wrong.result["error"]["message"]


def test_a_copied_listing_line_or_id_and_name_still_names_the_item(ws, index):
    """Measured: after files(list) the model passed 'data_1 shop_sales.csv' and every read failed."""
    ws.add_bytes("file", b"print(1)", "py", name="parabola.py")
    line = ws.lines()["file_1"]
    for ref in (line, "file_1 parabola.py", "parabola.py (file_1)", " 'PARABOLA.PY' "):
        assert run(ws, index, "read", name=ref, start=1).result["file"] == "file_1", ref
    assert run(ws, index, "read", name="parabola.py", start=0).result["lines"].startswith("1-")


def test_rename_to_the_same_name_is_refused(ws, index):
    ws.add_bytes("file", b"x", "rs", name="tts.rs")
    same = run(ws, index, "rename", name="tts.rs", new_name="tts.rs")
    assert same.status == "invalid_arguments" and "already called tts.rs" in same.result["error"]["message"]


def test_read_in_lines_with_the_next_step(ws, index):
    ws.add_bytes("file", "\n".join(f"x = {i}  # " + "padding " * 8 for i in range(400)).encode(), "py",
                 name="long.py")
    first = run(ws, index, "read", name="long.py").result  # main reads up to 8,000 characters at a time
    shown = int(first["lines"].split("-")[1].split()[0])
    assert first["text"].startswith("  1| x = 0") and 50 < shown < 400  # code comes numbered
    assert first["next"] == f"files(mode='read', name='long.py', start={shown + 1})"
    end = run(ws, index, "read", name="long.py", start=391).result
    assert end["lines"] == "391-400 of 400" and end["next"] == "end of file"
    table = ws.add_table(pd.DataFrame({"a": range(30)}), "t.csv")
    read = run(ws, index, "read", name=table).result
    assert read["rows"] == 30 and len(read["first_rows"]) == 10 and "analyze_or_code" in read["note"]
    assert run(ws, index, "read").result["error"]["message"] == "read needs name: a file's name, e.g. 'report.md'."


def test_write_new_files_by_kind(ws, index):
    md = run(ws, index, "write", name="report", content="# Sales\n\nRevenue grew in **March**.").result
    assert md["name"] == "report.md" and md["kind"] == "document"  # no suffix: Markdown, and searchable
    assert find_files(ws, index, "revenue grew").result["files"][0]["path"] == "report.md"
    csv = run(ws, index, "write", name="list.csv", content="item,qty\npens,4\n").result
    assert csv["kind"] == "table" and ws.table(csv["file"])["qty"].tolist() == [4]
    code = run(ws, index, "write", name="hello.rs", content='fn main() { println!("hi"); }')
    assert code.result["kind"] == "file" and code.artifact == {"type": "file", "id": code.result["file"],
                                                                "name": "hello.rs", "kind": "file", "bytes": 29,
                                                                "action": "written"}
    assert ws.entry(code.result["file"])["by"] == "assistant"
    assert run(ws, index, "write", name="../x.txt", content="x").status == "invalid_arguments"


def test_a_docx_is_made_from_markdown(ws, index):
    from docx import Document
    data = markdown_docx("# Title\n\nSome **bold** text.\n\n- one\n- two\n\n| a | b |\n|---|---|\n| 1 | 2 |")
    doc = Document(io.BytesIO(data))
    assert [p.text for p in doc.paragraphs][:4] == ["Title", "Some bold text.", "one", "two"]
    assert doc.paragraphs[1].runs[1].bold and doc.tables[0].cell(1, 1).text == "2"
    made = run(ws, index, "write", name="summary.docx", content="# Summary\n\nAll good.").result
    assert made["kind"] == "document" and find_files(ws, index, "all good").result["files"]


def test_write_makes_a_file_or_gives_one_its_whole_new_text(ws, index):
    run(ws, index, "write", name="notes.txt", content="first")
    mine = run(ws, index, "write", name="notes.txt", content="second")
    assert mine.status == "ok" and ws.path("file_1").read_text() == "second"  # the assistant's own: no asking
    assert mine.artifact["action"] == "updated" and len(ws.versions("file_1")) == 1  # the first text kept
    picture = ws.add_image(_png(), "logo.png")
    assert "an image" in run(ws, index, "write", name="logo.png", content="x").result["error"]["message"]
    assert ws.path(picture).read_bytes() == _png()
    assert run(ws, index, "write", name="notes.txt").result["error"]["field"] == "content"
    assert run(ws, index, "update", name="notes.txt", content="x").result["error"]["field"] == "mode"  # gone


def test_edit_changes_one_passage(ws, index):
    run(ws, index, "write", name="calc.py", content="a = 1\nb = 2\nc = 3\nprint(a + b + c)\n")
    run(ws, index, "edit", name="calc.py", old_text="b = 2\nc = 3", new_text="b = 20\nc = 30\nd = 40")
    assert ws.path("file_1").read_text() == "a = 1\nb = 20\nc = 30\nd = 40\nprint(a + b + c)\n"
    run(ws, index, "edit", name="calc.py", old_text="a = 1", new_text="a = 100")
    assert ws.path("file_1").read_text().splitlines()[0] == "a = 100" and len(ws.versions("file_1")) == 2
    run(ws, index, "write", name="plan.docx", content="# Plan")
    assert "written whole" in run(ws, index, "edit", name="plan.docx", old_text="Plan", new_text="x").result["error"]["message"]


def test_changing_the_users_own_file_asks_first(ws, index):
    upload = ws.add_bytes("file", b"original", "txt", name="upload.txt", by="user")
    ask = run(ws, index, "write", name="upload.txt", content="changed").artifact
    assert ask["options"] == ["Update", "Cancel"] and ws.path(upload).read_bytes() == b"original"
    approval.carry_out(ws, index, ask["approval"], ask["options"], "Cancel")
    assert ws.path(upload).read_bytes() == b"original"
    approval.carry_out(ws, index, ask["approval"], ask["options"], "Update")
    assert ws.path(upload).read_bytes() == b"changed" and len(ws.versions(upload)) == 1  # the original kept


def test_rename_keeps_the_suffix_and_refuses_a_taken_name(ws, index):
    run(ws, index, "write", name="draft.md", content="# Plan\n\nShip in May.")
    ws.add_bytes("file", b"x", "txt", name="other.txt")
    done = run(ws, index, "rename", name="draft.md", new_name="final")
    assert done.result["renamed"] == "draft.md -> final.md" and ws.entry("file_1")["name"] == "final.md"
    assert find_files(ws, index, "ship in may").result["files"][0]["path"] == "final.md"
    taken = run(ws, index, "rename", name="final.md", new_name="other.txt")
    assert taken.status == "invalid_arguments" and "already a file here" in taken.result["error"]["message"]
    assert run(ws, index, "rename", name="final.md").result["error"]["field"] == "new_name"


def test_reports_are_read_first_listed_and_searchable(ws, index):
    from app.tools.files import index_report
    ws.add_bytes("file", b"import numpy as np\nprint(np.pi)\n", "py", name="pi.py")
    ws.set_report("file_1", "# Prints pi with numpy\n\n## How it works\nOne call to np.pi; ran OK.")
    index_report(ws, index, "file_1")
    read = run(ws, index, "read", name="pi.py").result
    assert "How it works" in read["report"] and read["next"].endswith("start=1)")
    assert "np.pi" in run(ws, index, "read", name="pi.py", start=1).result["text"]  # the content, on request
    assert run(ws, index, "list").result["items"][0].endswith("| Prints pi with numpy")
    hit = find_files(ws, index, "how it works, one call").result["files"][0]  # words only its report has
    assert hit["path"] == "pi.py" and hit["match"].startswith("its report: ")
    ws.set_report("file_1", "# Prints pi to 3 places")  # the latest only
    assert ws.report("file_1") == "# Prints pi to 3 places" and ws.lines()["file_1"].endswith("| Prints pi to 3 places")
    report_file = ws.root / ws.entry("file_1")["report"]["file"]
    approval.delete(ws, index, ["file_1"])
    assert not report_file.exists() and not find_files(ws, index, "np.pi").result["files"]


def test_delete_and_download_ask_first(ws, index, monkeypatch):
    ws.add_bytes("file", b"x", "txt", name="old.txt")
    ask = run(ws, index, "delete", name="old.txt").artifact
    assert ask["approval"]["action"] == "delete_files" and ask["approval"]["items"] == ["file_1"]
    assert ws.ids() == ["file_1"]  # nothing deleted yet
    asked = []
    monkeypatch.setattr(approval, "download_request", lambda link: asked.append(("github", link)) or "github")
    monkeypatch.setattr(approval, "import_request", lambda link, table=None: asked.append(("table", link)) or "table")
    run(ws, index, "download", link="https://github.com/o/r/blob/main/app.py")
    run(ws, index, "download", link="https://data.example.org/sales.csv")
    assert asked == [("github", "https://github.com/o/r/blob/main/app.py"), ("table", "https://data.example.org/sales.csv")]
    assert run(ws, index, "download").result["error"]["field"] == "link"


def test_text_is_never_written_into_a_binary_format(ws, index):
    """Read in the code: write, update and rename put text bytes under any name, so 'report.pdf' or an
    uploaded PDF changed as text became a broken file under a good-looking name."""
    refused = run(ws, index, "write", name="report.pdf", content="# Report")
    assert refused.status == "invalid_arguments" and "isn't text" in refused.result["error"]["message"]
    assert not ws.ids()
    pdf = ws.add_bytes("file", b"%PDF-1.4 ...", "pdf", name="terms.pdf")
    assert run(ws, index, "update", name="terms.pdf", content="new text").status == "invalid_arguments"  # not asked
    note = run(ws, index, "write", name="notes.md", content="hi").result["file"]
    assert run(ws, index, "rename", name=note, new_name="notes.pdf").status == "invalid_arguments"
    assert run(ws, index, "write", name="plan.docx", content="# Plan").status == "ok"  # built from Markdown
    assert ws.path(pdf).read_bytes() == b"%PDF-1.4 ..."


def test_a_program_main_writes_is_parsed_and_said_to_be_unrun(ws, index):
    """Measured on the 26B: main fixed stats.py with files(update), left it broken, and said what it 'now prints'."""
    good = files(ws, index, "write", name="mean.py", content="print(sum([1, 2]) / 2)\n").result
    assert "not run" in good["not_run"] and "analyze_or_code" in good["not_run"] and "broken" not in good
    bad = files(ws, index, "write", name="mean.py", content="def mean(values):\ndef mean(values):\n    return 1\n").result
    assert bad["broken"].startswith("mean.py does not parse") and "line 2" in bad["broken"]
    assert "not_run" not in files(ws, index, "write", name="notes.md", content="# Notes\n").result
    # what the page does with it (measured: "run it in your terminal", "open it in your web browser", then the game
    # played in the chat by the model itself)
    assert "from its card in the chat (Run), and type into it" in good["not_run"]
    page = files(ws, index, "write", name="cups.html", content="<button>Cup 1</button>\n").result
    assert page["not_run"].startswith("cups.html works in the chat, under its card") and "analyze_or_code" not in page["not_run"]


def _png() -> bytes:
    from PIL import Image
    out = io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(out, "PNG")
    return out.getvalue()
