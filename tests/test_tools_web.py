"""web_search (fake engine) and read_webpage (fake server and DNS) - no network."""

from __future__ import annotations

import socket

import httpx
import numpy as np
import pytest

from app.rag.embed import EmbedderUnavailable
from app.tools import web
from app.tools.web import read_webpage, web_search

# ---------------------------------------------------------------- web_search


def engine(*answers):
    """A fake search engine: one answer (list of results, or an exception) per tier asked."""
    asked = []

    def search(query, k, tier, timeout, timelimit=None):
        asked.append((query, tier))
        search.timelimits.append(timelimit)
        answer = answers[len(asked) - 1]
        if isinstance(answer, Exception):
            raise answer
        return answer
    search.asked, search.timelimits = asked, []
    return search


HITS = [{"title": "Python 3.14 released", "href": "https://www.python.org/downloads/release/python-3140/",
         "body": "Python 3.14.0 is the newest major release."},
        {"title": "What's new", "href": "https://docs.python.org/3/whatsnew/3.14.html", "body": "Summary  of\nchanges."},
        {"title": "Dup", "href": "https://www.python.org/downloads/release/python-3140/", "body": "again"},
        {"title": "Elsewhere", "href": "https://example.com/python", "body": "unrelated site"}]


class Meaning:
    """A stand-in embedding model: a result means the query when its snippet holds ``word``."""
    name = "fake"

    def __init__(self, word, down=False):
        self.word, self.down = word, down

    def embed_query(self, text):
        return np.array([1.0, 0.0])

    def embed_documents(self, texts, titles=None):
        if self.down:
            raise EmbedderUnavailable("down")
        return np.array([[1.0, 0.0] if self.word in t else [0.0, 1.0] for t in texts])


FIVE = [{"title": f"Result {i}", "href": f"https://e{i}.example.com/", "body": "about cats" if i == 3 else "about dogs"}
        for i in range(5)]


def test_results_are_fused_with_meaning():
    order = [r["title"] for r in web_search("cats", engine=engine(FIVE), embedder=Meaning("cats")).result["results"]]
    assert order == ["Result 0", "Result 3", "Result 1", "Result 2", "Result 4"]  # up from 4th; the engines' first stays


def test_without_the_model_the_engines_order_stays():
    for embedder in (None, Meaning("cats", down=True)):
        order = [r["title"] for r in web_search("cats", engine=engine(FIVE), embedder=embedder).result["results"]]
        assert order == [f"Result {i}" for i in range(5)]


def test_results_are_cleaned_and_deduplicated():
    e = engine(HITS)
    r = web_search("python 3.14 release", engine=e)
    assert r.status == "ok" and len(r.result["results"]) == 3
    assert r.result["results"][1]["snippet"] == "Summary of changes."
    assert "read_webpage" in r.result["note"] and len(e.asked) == 1


def test_second_tier_only_when_the_first_finds_nothing():
    class DDGSException(Exception):
        pass
    e = engine(DDGSException("No results found."), HITS[:1])
    r = web_search("rare query", engine=e)
    assert [t for _, t in e.asked] == [e.asked[0][1], e.asked[1][1]] and len(e.asked) == 2
    assert r.result["results"][0]["title"] == "Python 3.14 released"
    r = web_search("nothing at all", engine=engine([], []))
    assert r.status == "ok" and r.result["results"] == [] and "No results" in r.result["note"]


def test_rate_limit_is_a_clear_failure():
    class RatelimitException(Exception):
        pass
    r = web_search("anything", engine=engine(RatelimitException("202 Ratelimit"), RatelimitException("x")))
    assert r.status == "failed" and "rate limiting" in r.result["error"]["message"]


def test_site_restricts_results():
    e = engine(HITS)
    r = web_search("whatsnew", site="https://www.python.org/", engine=e)
    assert e.asked[0][0] == "whatsnew site:python.org"
    assert {x["url"].split("/")[2] for x in r.result["results"]} == {"www.python.org", "docs.python.org"}
    assert web_search("x y", site="not a domain", engine=e).status == "invalid_arguments"


# ---------------------------------------------------------------- read_webpage

PUBLIC = {"news.example.com": "93.184.216.34", "anime.example.com": "93.184.216.40", "data.example.com": "93.184.216.41", "docs.example.com": "93.184.216.35", "wiki.example.com": "93.184.216.36",
          "sneaky.example.com": "10.0.0.5", "rebind.example.com": "127.0.0.1", "app.example.com": "93.184.216.37",
          "shop.example.com": "93.184.216.38"}


def resolve(host, port, proto=0):
    if host not in PUBLIC:
        try:
            socket.inet_aton(host)
            ip = host
        except OSError:
            raise socket.gaierror("no such host") from None
    else:
        ip = PUBLIC[host]
    return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (ip, port))]


ARTICLE = b"""<!doctype html><html><head><title>Rates rise again</title><script>var x = 'tracker';</script></head>
<body><nav><a href="/">Home</a> | <a href="/world">World</a></nav>
<header>Daily News - subscribe!</header>
<article><h1>Rates rise again</h1><p>The central bank raised rates by 0.25 points on Tuesday.</p>
<p>Markets had expected the move.</p><ul><li>Inflation: 3.1%</li><li>Next meeting: March</li></ul></article>
<aside>Related: celebrity gossip</aside><footer>(c) Daily News</footer>
<div hidden>secret tracking text</div></body></html>"""


def site(routes):
    def handler(request):
        answer = routes.get(str(request.url))
        if answer is None:
            return httpx.Response(404)
        return answer() if callable(answer) else answer
    return httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)


STREAMED_PAGE = b"""<!doctype html><html><head><title>Anime returns</title></head><body>
<main><h1>Anime returns</h1><!--$?--><template id="B:0"></template><p>Loading the story ...</p><!--/$-->
<p>Before the parts.</p><template id="P:1"></template><p>After the parts.</p></main>
<div hidden id="S:0"><p>Episode 1180 aired on September 27 as the last of the year; the series returns in 2027.</p></div>
<div hidden id="S:1"><p>Crunchyroll streams it in 84 territories.</p></div>
<div hidden>secret tracking text</div>
<script>$RC("B:0","S:0")</script><script>$RS("S:1","P:1")</script></body></html>"""


def test_a_streamed_page_reads_its_parts_where_they_belong():
    """React streams a page's parts late, hidden until its script moves them: they are read, in their place."""
    http = site({"https://anime.example.com/news": httpx.Response(200, content=STREAMED_PAGE,
                                                                   headers={"content-type": "text/html"})})
    text = read_webpage("https://anime.example.com/news", http=http, resolve=resolve).result["text"]
    assert "returns in 2027" in text and "84 territories" in text
    assert "Loading the story" not in text and "secret" not in text  # a boundary's fallback goes; other hidden text too
    assert text.index("Before the parts") < text.index("84 territories") < text.index("After the parts")


def test_article_text_without_page_furniture():
    http = site({"https://news.example.com/rates": httpx.Response(200, content=ARTICLE,
                                                                   headers={"content-type": "text/html; charset=utf-8"})})
    r = read_webpage("news.example.com/rates", http=http, resolve=resolve)  # https:// is added
    assert r.status == "ok" and r.result["title"] == "Rates rise again"
    text = r.result["text"]
    assert text.startswith("# Rates rise again") and "raised rates by 0.25 points" in text and "- Inflation: 3.1%" in text
    for junk in ("tracker", "Home", "subscribe", "gossip", "(c) Daily News", "secret"):
        assert junk not in text


def test_private_addresses_are_refused_even_after_a_redirect():
    http = site({"https://news.example.com/go": httpx.Response(302, headers={"location": "http://127.0.0.1:8080/admin"}),
                 "https://news.example.com/go2": httpx.Response(301, headers={"location": "https://sneaky.example.com/"})})
    for url in ["https://news.example.com/go", "https://news.example.com/go2", "http://localhost:8000/api",
                "http://192.168.1.1/", "https://rebind.example.com/", "http://[::1]/"]:
        r = read_webpage(url, http=http, resolve=resolve)
        assert r.status == "invalid_arguments" and "private network" in r.result["error"]["message"], url
    assert read_webpage("ftp://news.example.com/file", http=http, resolve=resolve).status == "invalid_arguments"


def test_redirects_between_public_pages_are_followed():
    http = site({"https://news.example.com/old": httpx.Response(301, headers={"location": "/new"}),
                 "https://news.example.com/new": httpx.Response(200, text="Plain text page.",
                                                                headers={"content-type": "text/plain"})})
    r = read_webpage("https://news.example.com/old", http=http, resolve=resolve)
    assert r.result["url"] == "https://news.example.com/new" and r.result["text"] == "Plain text page."


def test_large_binary_missing_and_empty_pages_fail_clearly():
    http = site({"https://docs.example.com/big": lambda: httpx.Response(200, content=b"a" * (3 * 1024 * 1024 + 10),
                                                                         headers={"content-type": "text/plain"}),
                 "https://docs.example.com/img": httpx.Response(200, content=b"\x89PNG....", headers={"content-type": "image/png"}),
                 "https://docs.example.com/app": httpx.Response(200, content=b"<html><body><div id=app></div><script>x</script></body></html>",
                                                                headers={"content-type": "text/html"})})
    assert "larger than 3 MB" in read_webpage("https://docs.example.com/big", http=http, resolve=resolve).result["error"]["message"]
    message = read_webpage("https://docs.example.com/img", http=http, resolve=resolve).result["error"]["message"]
    assert "an image" in message and "no text" in message
    assert "JavaScript" in read_webpage("https://docs.example.com/app", http=http, resolve=resolve).result["error"]["message"]
    r = read_webpage("https://docs.example.com/missing", http=http, resolve=resolve)
    assert r.status == "failed" and "404" in r.result["error"]["message"]
    assert read_webpage("https://nowhere.example.org/", http=http, resolve=resolve).status == "failed"


@pytest.mark.parametrize("charset_header", [True, False])
def test_page_encoding_is_respected(charset_header):
    page = "<html><head><meta charset='windows-1252'><title>Café</title></head><body><p>Crème brûlée</p></body></html>"
    headers = {"content-type": "text/html; charset=windows-1252" if charset_header else "text/html"}
    http = site({"https://news.example.com/menu": httpx.Response(200, content=page.encode("cp1252"), headers=headers)})
    r = read_webpage("https://news.example.com/menu", http=http, resolve=resolve)
    assert r.result["title"] == "Café" and "Crème brûlée" in r.result["text"]


# ---------------------------------------------------------------- metadata for the UI

CARD_PAGE = b"""<html lang="en-GB"><head><title>Rates rise | Daily News</title>
<meta property="og:site_name" content="Daily News"><meta property="og:image" content="/img/rates.jpg">
<meta name="description" content="The central bank raised rates again."><meta name="author" content="A. Writer">
<meta property="article:published_time" content="2026-09-30T08:00:00Z">
<link rel="shortcut icon" href="/static/icon.png"></head><body><article><p>Rates went up.</p></article></body></html>"""


def png(size=(16, 16)) -> bytes:
    import io
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", size, "red").save(buf, format="PNG")
    return buf.getvalue()


def test_page_card_comes_from_the_head():
    http = site({"https://news.example.com/rates": httpx.Response(200, content=CARD_PAGE,
                                                                   headers={"content-type": "text/html"})})
    r = read_webpage("https://news.example.com/rates", http=http, resolve=resolve)
    card = r.artifact
    assert card["site"] == "Daily News" and card["image"] == "https://news.example.com/img/rates.jpg"
    assert card["favicon"] == "https://news.example.com/static/icon.png" and card["language"] == "en"
    assert card["published"] == "2026-09-30T08:00:00Z" and card["author"] == "A. Writer"
    assert card["description"] == "The central bank raised rates again."
    assert r.result["published"] == "2026-09-30T08:00:00Z" and "image" not in r.result  # the model gets the essentials


def test_search_results_carry_site_and_favicon():
    r = web_search("x y", engine=engine(HITS[:1]))
    assert r.result["results"][0]["site"] == "python.org"
    assert r.result["results"][0]["favicon"] == "https://www.python.org/favicon.ico"


def test_icons_are_fetched_checked_and_cached():
    from app.tools import web
    web._icons.clear()
    calls = []
    http = site({"https://news.example.com/static/icon.png": lambda: (calls.append(1), httpx.Response(200, content=png()))[1],
                 "https://docs.example.com/favicon.ico": httpx.Response(200, content=b"<html>not an icon</html>")})
    icon = web.fetch_icon("https://news.example.com/rates", "https://news.example.com/static/icon.png", http=http, resolve=resolve)
    assert icon and icon[1] == "image/png"
    assert web.fetch_icon("https://news.example.com/other", None, http=http, resolve=resolve) == icon and len(calls) == 1
    assert web.fetch_icon("https://docs.example.com/page", None, http=http, resolve=resolve) is None  # HTML is not an icon
    # an icon link pointing into a private network is never fetched
    assert web.fetch_icon("https://sneaky.example.com/", "http://127.0.0.1/icon.png", http=http, resolve=resolve) is None


def test_preview_images_must_be_real_raster_images():
    from app.tools import web
    svg = b"<svg xmlns='http://www.w3.org/2000/svg'><script>alert(1)</script></svg>"
    http = site({"https://news.example.com/a.png": httpx.Response(200, content=png((64, 32))),
                 "https://news.example.com/b.svg": httpx.Response(200, content=svg, headers={"content-type": "image/svg+xml"})})
    assert web.fetch_image("https://news.example.com/a.png", http=http, resolve=resolve)[1] == "image/png"
    assert web.fetch_image("https://news.example.com/b.svg", http=http, resolve=resolve) is None


def test_long_pages_can_be_searched_with_find():
    filler = "\n".join(f"<p>Option number {i} does something unrelated to networking at all.</p>" for i in range(400))
    page = f"<html><body><main><h1>Server</h1><p>Intro.</p>{filler}<p>--port N: port to listen (default: 8080)</p></main></body></html>"
    http = site({"https://docs.example.com/server": lambda: httpx.Response(200, content=page.encode(),
                                                                            headers={"content-type": "text/html"})})
    cut = read_webpage("https://docs.example.com/server", http=http, resolve=resolve).result
    assert "8080" not in cut["text"] and "find=" in cut["truncated"]
    found = read_webpage("https://docs.example.com/server", "default port", http=http, resolve=resolve).result
    assert "default: 8080" in found["text"] and found["text"].startswith("# Server") and "parts mentioning" in found["shown"]


def test_find_prefers_lines_with_more_of_the_words():
    noise = "\n".join(f"<p>--opt-{i} N: default value {i} (support for imports)</p>" for i in range(600))
    page = f"<html><body><main><h1>Server</h1>{noise}<p>--port N: port to listen (default: 8080)</p></main></body></html>"
    http = site({"https://docs.example.com/server": lambda: httpx.Response(200, content=page.encode(),
                                                                            headers={"content-type": "text/html"})})
    found = read_webpage("https://docs.example.com/server", "default port", http=http, resolve=resolve).result
    assert "default: 8080" in found["text"]  # one line with both words beats 600 with one ('support' isn't 'port')


def test_failed_pages_say_what_to_do_next():
    http = site({"https://docs.example.com/busy": httpx.Response(503)})
    message = read_webpage("https://docs.example.com/busy", http=http, resolve=resolve).result["error"]["message"]
    assert "503" in message and "read another source" in message


SERVICES = {"github.com": "140.82.112.4", "api.github.com": "140.82.112.6", "raw.githubusercontent.com": "185.199.108.133",
            "www.youtube.com": "142.250.183.14"}


def services_resolve(host, port, proto=0):
    if host in SERVICES:
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (SERVICES[host], port))]
    return resolve(host, port, proto)


def test_github_pages_are_read_through_its_api():
    repo = {"full_name": "acme/db", "description": "A fast database.", "language": "C++",
            "license": {"spdx_id": "MIT"}, "stargazers_count": 1234}
    issue = b"<html><head><title>Crash on start</title></head><body><main><p>Crash on start when the port is taken.</p></main></body></html>"
    routes = {"https://api.github.com/repos/acme/db": httpx.Response(200, json=repo),
              "https://api.github.com/repos/acme/db/readme": httpx.Response(200, text="# db\nRun it on port 9000."),
              "https://api.github.com/repos/acme/db/contents/src/server.py?ref=main": httpx.Response(403, text="rate limit"),
              "https://raw.githubusercontent.com/acme/db/main/src/server.py": httpx.Response(200, text="PORT = 9000\n"),
              "https://github.com/acme/db/issues/7": httpx.Response(200, content=issue, headers={"content-type": "text/html"})}
    asked = []

    def handler(request):
        asked.append((str(request.url), request.headers.get("accept")))
        return routes.get(str(request.url), httpx.Response(404))
    http = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)

    r = read_webpage("https://github.com/acme/db", http=http, resolve=services_resolve)
    assert r.status == "ok" and r.result["url"] == "https://github.com/acme/db" and r.result["title"] == "acme/db"
    assert "Main language: C++. License: MIT." in r.result["text"] and "Stars: 1,234." in r.result["text"]
    assert "port 9000" in r.result["text"] and r.artifact["description"] == "A fast database."
    assert ("https://api.github.com/repos/acme/db/readme", "application/vnd.github.raw+json") in asked
    f = read_webpage("https://github.com/acme/db/blob/main/src/server.py", http=http, resolve=services_resolve)
    assert f.result["text"] == "PORT = 9000" and f.result["title"] == "acme/db: src/server.py"  # API refused -> raw file
    other = read_webpage("https://github.com/acme/db/issues/7", http=http, resolve=services_resolve)
    assert other.status == "ok" and "port is taken" in other.result["text"]  # other GitHub pages: read as HTML


WATCH = r"""<html><head><meta property="og:title" content="How DuckDB Works">
<meta property="og:description" content="Short version."></head><body>
<div itemscope itemtype="http://schema.org/VideoObject"><meta itemprop="duration" content="PT1H2M5S">
<meta itemprop="uploadDate" content="2024-05-01T08:00:00-07:00"><meta itemprop="genre" content="Education">
<span itemprop="author" itemscope itemtype="http://schema.org/Person"><link itemprop="name" content="Data Talks"></span></div>
<script>var ytInitialPlayerResponse = {"videoDetails":{"viewCount":"12345","shortDescription":"Chapters:\n0:00 Intro\n5:10 \"Vectors\""}};</script>
<footer>(c) 2026 Google LLC</footer></body></html>"""


def test_youtube_videos_are_read_from_their_watch_page():
    asked = []

    def handler(request):
        asked.append(str(request.url))
        if request.url.path == "/watch" and request.url.params["v"] == "abcDEF12345":
            return httpx.Response(200, text=WATCH, headers={"content-type": "text/html"})
        if request.url.path == "/watch":  # a consent page instead of the video
            return httpx.Response(200, text="<html><body><p>Before you continue</p></body></html>")
        if request.url.path == "/oembed":
            return httpx.Response(200, json={"title": "A short", "author_name": "Someone"})
        return httpx.Response(404)
    http = httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False)

    r = read_webpage("https://youtu.be/abcDEF12345", http=http, resolve=services_resolve)
    text = r.result["text"]
    assert r.status == "ok" and r.result["title"] == "How DuckDB Works" and r.result["author"] == "Data Talks"
    assert "Length: 1:02:05. Published: 2024-05-01. Views: 12,345. Category: Education." in text
    assert '5:10 "Vectors"' in text and "can't be read" in text and "Google LLC" not in text
    assert r.artifact["video"] == {"id": "abcDEF12345", "embed": "https://www.youtube-nocookie.com/embed/abcDEF12345",
                                   "length": "1:02:05"}
    assert r.artifact["image"] == "https://i.ytimg.com/vi/abcDEF12345/hqdefault.jpg"
    s = read_webpage("https://www.youtube.com/shorts/zzzZZZ00000", http=http, resolve=services_resolve)
    assert s.result["title"] == "A short" and "Channel: Someone." in s.result["text"]  # oEmbed fallback
    assert web.youtube_id("https://www.youtube.com/@channel") is None
    assert web.youtube_id("https://m.youtube.com/watch?v=abcDEF12345&t=30") == "abcDEF12345"


def test_video_search_results_carry_a_thumbnail_for_the_ui():
    hits = [{"title": "Talk", "href": "https://www.youtube.com/watch?v=abcDEF12345", "body": "a talk"}, HITS[0]]
    r = web_search("duckdb talk", engine=engine(hits))
    assert r.artifact["items"][0]["image"] == "https://i.ytimg.com/vi/abcDEF12345/hqdefault.jpg"
    assert "image" not in r.result["results"][0] and "image" not in r.artifact["items"][1]  # the model doesn't need it


def test_pages_drawn_by_javascript_give_their_description():
    page = b'<html><head><meta name="description" content="Fast analytics for everyone."></head><body><div id=app></div></body></html>'
    http = site({"https://docs.example.com/spa": httpx.Response(200, content=page, headers={"content-type": "text/html"})})
    r = read_webpage("https://docs.example.com/spa", http=http, resolve=resolve)
    assert r.status == "ok" and r.result["text"].startswith("Fast analytics for everyone.") and "JavaScript" in r.result["text"]


def test_a_bot_check_page_is_reported_not_read():
    from app.tools.web import _bot_check
    seen = "##### Verifying your browser…\nThis takes less than a second.\nYour experience on this site will be improved."
    assert _bot_check("Verifying… — Danielitte", seen, "https://danielitte.com/recaptcha/challenge")
    assert _bot_check("Just a moment...", "Enable JavaScript and cookies to continue", "https://x.org/a")
    assert not _bot_check("Short note", "Our newsletter. Subscribe for updates.", "https://x.org/b")
    article = "A CAPTCHA is a type of challenge-response test. " * 40  # a page about captchas is read
    assert not _bot_check("CAPTCHA - Wikipedia", article, "https://en.wikipedia.org/wiki/CAPTCHA")


# ---------------------------------------------------------------- page age and the files a page offers

DATASETS = b"""<!doctype html><html><head><title>Open data: city bikes</title>
<meta property="article:published_time" content="2023-05-01T09:00:00Z">
<meta property="article:modified_time" content="2026-03-02">
</head><body><article><h1>City bikes</h1><p>Trips per station, updated monthly. Download the data below.</p>
<p><a href="/files/trips-2025.csv">Trips 2025 (CSV)</a> <a href="/files/stations.xlsx">Stations</a>
<a href="/files/readme.txt">notes</a> <a href="/about">About us</a>
<a href="https://github.com/city/bikes/blob/main/clean.py">cleaning script</a>
<a href="https://github.com/city/bikes">the repository</a></p></article></body></html>"""


def test_a_page_says_how_old_it_is_and_which_files_it_offers():
    from datetime import date
    assert web.page_age("2023-05-01T09:00:00Z", None, today=date(2026, 10, 5)) == "published 2023-05-01, 3 years ago"
    assert web.page_age("2023-05-01", "2026-09-20", today=date(2026, 10, 5)) == "updated 2026-09-20, 15 days ago"
    assert web.page_age("May 2023", None) is None
    http = site({"https://news.example.com/bikes": httpx.Response(200, content=DATASETS,
                                                                  headers={"content-type": "text/html"})})
    r = read_webpage("https://news.example.com/bikes", http=http, resolve=resolve).result
    assert r["updated"] == "2026-03-02" and r["age"].startswith("updated 2026-03-02, ")
    assert r["files"] == ["https://news.example.com/files/trips-2025.csv", "https://news.example.com/files/stations.xlsx",
                          "https://github.com/city/bikes/blob/main/clean.py"]  # tables and GitHub files only
    assert "give its link in downloads" in r["files_note"]


# ---------------------------------------------------------------- statuses, wiki pages, recent searches

def test_recent_searches_use_the_engines_time_filter():
    e = engine(HITS)
    r = web_search("one piece latest episode", recent="week", engine=e)
    assert r.status == "ok" and r.result["recent"] == "week" and e.timelimits == ["w"]
    assert web_search("x y", recent="decade", engine=e).status == "invalid_arguments"
    plain = engine(HITS)
    web_search("x y", engine=plain)
    assert plain.timelimits == [None]


def test_every_read_says_how_it_went():
    long_page = b"<html><body><article>" + b"<p>" + b"word " * 4000 + b"</p></article></body></html>"
    http = site({"https://news.example.com/rates": httpx.Response(200, content=ARTICLE, headers={"content-type": "text/html"}),
                 "https://news.example.com/long": httpx.Response(200, content=long_page, headers={"content-type": "text/html"}),
                 "https://news.example.com/gone": httpx.Response(404),
                 "https://news.example.com/locked": httpx.Response(403)})
    assert read_webpage("https://news.example.com/rates", http=http, resolve=resolve).result["read"] == "ok"
    assert read_webpage("https://news.example.com/long", http=http, resolve=resolve).result["read"] == "cut"
    gone = read_webpage("https://news.example.com/gone", http=http, resolve=resolve).result["error"]
    locked = read_webpage("https://news.example.com/locked", http=http, resolve=resolve).result["error"]
    assert gone["reason"] == "not_found" and locked["reason"] == "blocked" and "another source" in locked["message"]
    nowhere = read_webpage("https://nowhere.example.com/", http=http, resolve=resolve).result["error"]
    assert nowhere["reason"] == "not_found"


WIKI_JSON = {"parse": {"title": "Episode 1071", "pageid": 1, "text":
             '<div class="mw-parser-output"><aside class="portable-infobox"><h2>Episode 1071</h2>'
             '<div>Airdate: August 6, 2023</div></aside><p>Luffy awakens Gear 5 against Kaido.</p></div>'}}


def test_a_wiki_page_that_blocks_readers_is_read_through_its_api():
    """Measured: Fandom answers its pages with Cloudflare's 'Just a moment' (403), and its api.php with the page."""
    import json as _json
    challenge = httpx.Response(403, content=b"<html><title>Just a moment...</title></html>", headers={"content-type": "text/html"})
    api = ("https://wiki.example.com/api.php?action=parse&page=Episode_1071&prop=text|displaytitle&format=json"
           "&formatversion=2&redirects=1&disableeditsection=1")
    http = site({"https://wiki.example.com/wiki/Episode_1071": challenge,
                 api: httpx.Response(200, content=_json.dumps(WIKI_JSON).encode(), headers={"content-type": "application/json"}),
                 "https://wiki.example.com/wiki/Missing": challenge,
                 api.replace("Episode_1071", "Missing"): httpx.Response(200, json={"error": {"code": "missingtitle"}}),
                 "https://news.example.com/wiki/Thing": challenge})
    r = read_webpage("https://wiki.example.com/wiki/Episode_1071", http=http, resolve=resolve)
    assert r.status == "ok" and r.result["title"] == "Episode 1071" and r.result["via"] == "the wiki's API"
    assert "Gear 5" in r.result["text"] and "Airdate: August 6, 2023" in r.result["text"]  # the infobox is kept
    missing = read_webpage("https://wiki.example.com/wiki/Missing", http=http, resolve=resolve).result["error"]
    assert missing["reason"] == "not_found" and "no page called 'Missing'" in missing["message"]
    no_api = read_webpage("https://news.example.com/wiki/Thing", http=http, resolve=resolve).result["error"]
    assert no_api["reason"] == "blocked"  # no wiki API there: the page's own refusal stands


# ---------------------------------------------------------------- one page per result, pacing, the page cache

def test_the_same_page_twice_in_the_results_is_one_result():
    hits = [{"title": "A", "href": "https://www.example.com/post/?utm_source=feed", "body": "a"},
            {"title": "A again", "href": "https://example.com/post#comments", "body": "a"},
            {"title": "B", "href": "https://example.com/post?id=2", "body": "b"}]
    r = web_search("post", engine=engine(hits))
    assert [x["title"] for x in r.result["results"]] == ["A", "B"]


def test_a_page_read_again_within_the_hour_comes_from_the_cache(monkeypatch):
    calls = []
    monkeypatch.setattr(web, "_pages", {})
    monkeypatch.setattr(web, "_page", lambda http, url, resolve, browser=None: calls.append(url) or (
        "Title", "Some text on the page.", {"site": "example.com", "files": ["https://example.com/a.csv"]}, url))
    first = read_webpage("https://example.com/cached")
    again = read_webpage("https://example.com/cached")
    assert calls == ["https://example.com/cached"] and first.result == again.result
    assert again.result["files"] == ["https://example.com/a.csv"]  # the cached card isn't changed by a read


def test_pages_of_one_site_are_read_a_moment_apart(monkeypatch):
    waits = []
    monkeypatch.setattr(web, "_last_read", {})
    monkeypatch.setattr(web.time, "sleep", waits.append)
    web._pace("example.com")
    web._pace("example.com")
    web._pace("other.example")
    assert len(waits) == 1 and 0 < waits[0] <= web.SITE_GAP


# ---------------------------------------------------------------- the browser, only when the plain read can't

DRAWN = b"""<!doctype html><html><head><title>App</title></head><body><div id="root"></div>
<script src="/app.js"></script></body></html>"""
RENDERED = """<html><head><title>Pricing</title></head><body><main><h1>Pricing</h1>
<p>The basic plan costs 9 dollars a month and includes five projects with unlimited collaborators.</p>
<p>The team plan costs 29 dollars a month and adds single sign-on, an audit log and priority support.</p>
<my-card>Enterprise plans are priced per seat; ask the sales team for a quote that fits your company.</my-card>
</main></body></html>"""


def browser_showing(html, final=None, seen=None):
    def page(url, check):
        check(url)
        (seen if seen is not None else []).append(url)
        check(final or url)
        return final or url, "From the browser", html
    return page


def test_a_page_drawn_by_javascript_is_read_through_the_browser():
    http = site({"https://app.example.com/pricing": httpx.Response(200, content=DRAWN,
                                                                     headers={"content-type": "text/html"})})
    seen = []
    r = read_webpage("https://app.example.com/pricing", http=http, resolve=resolve,
                     browser=browser_showing(RENDERED, seen=seen))
    assert seen == ["https://app.example.com/pricing"] and r.result["via"] == "a browser"
    text = r.result["text"]
    assert "basic plan costs 9 dollars" in text and "priced per seat" in text  # a custom element's text counts too
    assert r.result["title"] == "Pricing" and r.result["read"] == "ok"


def test_a_page_the_plain_read_gets_never_opens_the_browser():
    http = site({"https://news.example.com/rates": httpx.Response(200, content=ARTICLE,
                                                                   headers={"content-type": "text/html"})})
    seen = []
    r = read_webpage("https://news.example.com/rates", http=http, resolve=resolve, browser=browser_showing("x", seen=seen))
    assert r.status == "ok" and seen == [] and "via" not in r.result


def test_a_refused_page_is_tried_in_the_browser_and_a_bot_check_there_is_still_blocked():
    http = site({"https://shop.example.com/item": httpx.Response(403)})
    r = read_webpage("https://shop.example.com/item", http=http, resolve=resolve, browser=browser_showing(RENDERED))
    assert r.status == "ok" and r.result["via"] == "a browser"
    check = "<html><head><title>Just a moment...</title></head><body><p>Verifying you are human.</p></body></html>"
    r = read_webpage("https://shop.example.com/item", http=http, resolve=resolve, browser=browser_showing(check))
    assert r.status == "failed" and r.result["error"]["reason"] == "blocked" and "even to a browser" in r.result["error"]["message"]


def test_the_browser_never_reads_a_page_on_this_computer():
    http = site({"https://shop.example.com/item": httpx.Response(403)})
    r = read_webpage("https://shop.example.com/item", http=http, resolve=resolve,
                     browser=browser_showing(RENDERED, final="http://127.0.0.1:8080/props"))  # it was sent there
    assert r.status == "invalid_arguments" and "this computer" in r.result["error"]["message"]


def test_without_a_working_browser_the_plain_outcome_stands():
    from app.tools.browser import McpError

    def broken(url, check):
        raise McpError("node is missing")
    http = site({"https://shop.example.com/item": httpx.Response(403)})
    r = read_webpage("https://shop.example.com/item", http=http, resolve=resolve, browser=broken)
    assert r.status == "failed" and r.result["error"]["reason"] == "blocked"


def test_a_cookie_notice_is_not_page_text():
    page = (b"<html><body><main><div>This site uses cookies to improve your experience. Learn more. OK, Got it.</div>"
            b"<p>Grandma bakes cookies every Sunday for the whole family, with butter and a lot of patience.</p>"
            b"<p>The real content is here.</p></main></body></html>")
    text = web._html_text(page, "utf-8", "https://news.example.com/x")[1]
    assert "uses cookies" not in text and "Grandma bakes cookies" in text and "real content" in text


# ---------------------------------------------------------------- links that aren't pages


def _office(kind: str) -> bytes:
    import io as _io
    out = _io.BytesIO()
    if kind == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading("Refund policy", 1)
        doc.add_paragraph("Opened phones can be returned within 14 days.")
        doc.save(out)
    else:
        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Q4 results"
        slide.placeholders[1].text = "Revenue grew 12% in the North"
        deck.save(out)
    return out.getvalue()


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def test_word_and_powerpoint_links_are_read_like_documents():
    http = site({"https://docs.example.com/policy": httpx.Response(200, content=_office("docx"), headers={"content-type": DOCX}),
                 "https://docs.example.com/deck.pptx": httpx.Response(200, content=_office("pptx"),
                                                                       headers={"content-type": "application/octet-stream"})})
    word = read_webpage("https://docs.example.com/policy", http=http, resolve=resolve).result
    assert word["read"] == "ok" and "## Refund policy" in word["text"] and "within 14 days" in word["text"]
    slides = read_webpage("https://docs.example.com/deck.pptx", http=http, resolve=resolve).result
    assert "[slide 1]" in slides["text"] and "## Q4 results" in slides["text"] and "grew 12%" in slides["text"]


def test_a_data_file_link_shows_its_first_rows_and_how_to_bring_it_in():
    rows = "region,sales\n" + "".join(f"R{i},{i}\n" for i in range(100))
    http = site({"https://data.example.com/sales.csv": httpx.Response(200, content=rows.encode(),
                                                                       headers={"content-type": "text/csv"})})
    r = read_webpage("https://data.example.com/sales.csv", http=http, resolve=resolve).result
    assert r["read"] == "ok" and "Columns (2): region, sales" in r["text"] and "R19,19" in r["text"]
    assert "R20,20" not in r["text"]  # the first rows only
    assert "data file" in r["note"] and "downloads" in r["note"] and "brings it in as a table" in r["note"] and "import_table" not in r["note"]


def test_an_excel_link_shows_its_first_rows():
    import io as _io

    import pandas as pd
    out = _io.BytesIO()
    pd.DataFrame({"city": ["Pune", "Delhi"], "temp": [31, 35]}).to_excel(out, index=False)
    http = site({"https://data.example.com/t.xlsx": httpx.Response(200, content=out.getvalue(), headers={
        "content-type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})})
    r = read_webpage("https://data.example.com/t.xlsx", http=http, resolve=resolve).result
    assert "city, temp" in r["text"] and "Pune,31" in r["text"] and "XLSX" in r["note"]
