"""Finished documents: every chart is a chart file whose picture is drawn when a document needs it (chart_image.py),
and Word and slide files are written from Markdown with those pictures in them (file_ops.py) - opened again with
python-docx and python-pptx."""

from __future__ import annotations

import io

import numpy as np
import pandas as pd
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from app.rag.embed import HashEmbedder
from app.tools import chart_image
from app.tools.file_ops import files, markdown_docx
from app.tools.files import search_files
from app.tools.find_patterns import find_patterns
from app.tools.forecast import forecast
from app.tools.make_chart import make_chart
from app.tools.search_index import SearchIndex
from app.tools.workspace import Workspace


@pytest.fixture
def ws(tmp_path):
    return Workspace(tmp_path / "ws")


@pytest.fixture
def index(tmp_path):
    return SearchIndex(tmp_path / "index.sqlite", HashEmbedder())


@pytest.fixture
def sales(ws):
    rng = np.random.default_rng(0)
    days = pd.date_range("2023-01-01", periods=730, freq="D")
    df = pd.DataFrame({"date": days, "region": rng.choice(["North", "South", "East"], len(days)),
                       "amount": rng.gamma(2.0, 50.0, len(days)).round(2), "units": rng.integers(1, 9, len(days))})
    return ws.add_table(df, "sales.csv")


def _picture(ws, item):
    with Image.open(chart_image.picture_of(ws, item)) as img:
        return img.format, img.size


@pytest.mark.parametrize("kind, args", [
    ("line", {"x": "date", "y": "amount"}), ("line", {"x": "date", "y": "amount", "by": "region"}),
    ("bar", {"x": "region", "y": "amount"}), ("bar", {"x": "region", "y": "amount", "by": "units", "stacked": True}),
    ("scatter", {"x": "units", "y": "amount", "by": "region"}), ("histogram", {"x": "amount"}),
    ("box", {"x": "region", "y": "amount"}), ("heatmap", {"x": "region", "by": "units"}), ("pie", {"x": "region", "y": "amount"})])
def test_every_chart_is_a_chart_file_that_draws_as_a_picture(ws, sales, kind, args):
    r = make_chart(ws, sales, kind, **args)
    key = r.result["chart_file"]
    assert r.status == "ok" and ws.kind(key) == "file" and ws.label(key).endswith(".chart.json")
    assert _picture(ws, key) == ("PNG", (960, 540)) and ws.entry(key)["parent"] == sales
    assert ws.ids("image") == []  # no picture copy in the workspace


def test_forecasts_and_unusual_periods_are_chart_files_too(ws, sales):
    assert _picture(ws, forecast(ws, sales, value="amount", time_unit="month").result["chart_file"])[0] == "PNG"
    assert _picture(ws, forecast(ws, sales, value="amount", time_unit="month", by="region").result["chart_file"])[0] == "PNG"
    assert _picture(ws, find_patterns(ws, sales, "unusual", value="amount").result["chart_file"])[0] == "PNG"


def test_a_chart_that_cannot_be_drawn_has_no_picture_and_still_shows(ws, sales, monkeypatch):
    key = make_chart(ws, sales, "bar", x="region", y="amount").result["chart_file"]
    monkeypatch.setattr(chart_image, "draw", lambda a: 1 / 0)
    assert chart_image.picture_of(ws, key) is None and chart_image.spec_of(ws, key)["chart"] == "bar"


def test_a_word_document_includes_a_chart(ws, index, sales):
    chart = ws.label(make_chart(ws, sales, "bar", x="region", y="amount").result["chart_file"])
    r = files(ws, index, "write", name="report.docx",
              content=f"# Sales report\n\nThe **North** leads.\n\n![Sales by region]({chart})\n")
    assert r.status == "ok"
    doc = Document(io.BytesIO(ws.path(r.result["file"]).read_bytes()))
    assert len(doc.inline_shapes) == 1 and any(p.text == "Sales by region" for p in doc.paragraphs)


def test_a_document_naming_a_picture_that_is_not_there_is_not_written(ws, index, sales):
    """Not saved with '[picture ... not found]' in it under a result that says 'Saved': refused, with the
    pictures there are, so the next call names one of them."""
    chart = ws.label(make_chart(ws, sales, "bar", x="region", y="amount").result["chart_file"])
    before = set(ws.ids())
    r = files(ws, index, "write", name="deck.pptx", content=f"# Sales\n![Sales]({chart})\n# More\n![Trend](image_99)\n")
    message = r.result["error"]["message"]
    assert r.status == "invalid_arguments" and "image_99" in message and chart in message
    assert set(ws.ids()) == before  # nothing written
    kept = files(ws, index, "write", name="notes.docx", content="# Notes\n- one\n").result["file"]
    assert files(ws, index, "write", name=kept, content="# Notes\n![x](chart_9.png)\n").status == "invalid_arguments"
    edited = Document(io.BytesIO(markdown_docx("![x](image_99)")))  # the user's own edit (the UI) shows where it was
    assert any("[picture image_99 not found]" in p.text for p in edited.paragraphs)


def test_slides_are_made_from_markdown_with_charts_and_notes(ws, index, sales):
    chart = make_chart(ws, sales, "line", x="date", y="amount").result["chart_file"]
    deck = (f"# Q4 sales review\nPrepared by Zen\n\n"
            f"# Highlights\n- Revenue grew **12%**\n  - mostly in the North\n- Fewer returns\n"
            f"![Revenue by month]({ws.label(chart)})\nNotes: start with the good news\n\n"
            f"# The trend\n![Revenue by month]({ws.entry(chart)['name']})\n\n## Next steps\n1. Hire\n2. Expand\n")
    r = files(ws, index, "write", name="review.pptx", content=deck)
    assert r.status == "ok" and r.result["kind"] == "document"  # slides are a document: searchable, like Word
    assert search_files(index, "revenue grew in the North").result["passages"]
    prs = Presentation(io.BytesIO(ws.path(r.result["file"]).read_bytes()))
    slides = list(prs.slides)
    assert [s.shapes.title.text for s in slides] == ["Q4 sales review", "Highlights", "The trend", "Next steps"]
    assert slides[0].placeholders[1].text == "Prepared by Zen"  # a title slide with its subtitle
    pictures = [sum(sh.shape_type == 13 for sh in s.shapes) for s in slides]  # 13: a picture
    assert pictures == [0, 1, 1, 0]
    text = " ".join(sh.text_frame.text for sh in slides[1].shapes if sh.has_text_frame)
    assert "Revenue grew 12%" in text and "mostly in the North" in text
    bold = [run for sh in slides[1].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for run in p.runs
            if run.text == "12%"]
    assert bold and bold[0].font.bold
    assert slides[1].notes_slide.notes_text_frame.text == "start with the good news"
    assert all(sh.left + sh.width <= prs.slide_width for s in slides for sh in s.shapes)  # nothing off the slide


def test_slides_need_a_heading_and_are_changed_whole(ws, index):
    assert files(ws, index, "write", name="empty.pptx", content="\n\n").status == "invalid_arguments"
    r = files(ws, index, "write", name="deck.pptx", content="# One\n- a\n")
    edit = files(ws, index, "edit", name="deck.pptx", old_text="- a", new_text="- b")
    assert edit.status == "invalid_arguments" and "written whole from Markdown" in edit.result["error"]["message"]
    assert files(ws, index, "write", name="deck.pptx", content="# One\n- b\n# Two\n").status == "ok"
    assert len(Presentation(io.BytesIO(ws.path(r.result["file"]).read_bytes())).slides) == 2
