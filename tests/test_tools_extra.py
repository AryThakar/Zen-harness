# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0
"""calculate, read_file, the workspace listing and the sandbox path() helper."""

from __future__ import annotations

import pandas as pd
import pytest

from app.tools.calc import calculate
from app.tools import agent_md
from app.tools.read_file import HELPER_CHARS, read_file
from app.tools.sandbox import run_python
from app.tools.workspace import Workspace


@pytest.fixture
def ws(tmp_path):
    return Workspace(tmp_path / "ws")


@pytest.mark.parametrize("expr, value", [("2400 * 0.15", 360.0), ("2 ^ 10", 1024), ("1,200 * 2", 2400),
                                         ("round((5240000 - 4810000) / 4810000 * 100, 2)", 8.94),
                                         ("sqrt(16) + log10(1000)", 7.0)])
def test_calculate(expr, value):
    r = calculate(expr)
    assert r.status == "ok" and r.result["result"] == pytest.approx(value)


@pytest.mark.parametrize("expr", ["1 / 0", "__import__('os').getcwd()", "foo(3)", "'a' * 3", "", "9 ** 9 ** 9"])
def test_calculate_refuses(expr):
    assert calculate(expr).status == "invalid_arguments"


def test_calculate_works_out_several_at_once():
    """Measured 2026-10-17: ten ratios in one expression were refused ('Tuple is not available'), and main worked
    them out in its head instead."""
    r = calculate("136.7 * 1000 / 211.9, 141.3 * 1000 / 216.8, round(2/3, 2)")
    assert r.result["results"] == pytest.approx([645.1156, 651.7528, 0.67], rel=1e-5)
    assert "In '1/0': Division by zero" in calculate("1/0, 2 + 2").result["error"]["message"]


# Written by the real model in bench/main/calc_probe (its own notation).
@pytest.mark.parametrize("expr, value", [("10!", 3628800), ("12! / (3! * 9!)", 220), ("(2+3)!", 120),
                                         ("sin(30 degrees)", 0.5), ("tan(35°)$", 0.7002075382),
                                         ("tan(35\\u00b0)", 0.7002075382), ("asin(0.5) * 180 / pi", 30),
                                         ("ln(50)", 3.9120230054), ("log2(1024)", 10), ("e^-0.5", 0.6065306597),
                                         ("2197^(1/3)", 13), ("(pi * ln(2)) / 8", 0.2721982613),
                                         # chat 7a918668b138: a catenary cable, refused before
                                         ("sinh(2)", 3.6268604078),
                                         ("5 * cosh(asinh((12 - 5 * sinh(2)) / 5))", 7.9138904155)])
def test_calculate_takes_the_models_notation(expr, value):
    r = calculate(expr)
    assert r.status == "ok" and r.result["result"] == pytest.approx(value) and "note" not in r.result


def test_a_plain_angle_is_said_to_be_radians():
    r = calculate("tan(35)")  # meant as degrees in the probe: the note lets the model correct it
    assert r.result["result"] == pytest.approx(0.4738147204) and "35°" in r.result["note"]
    assert "note" not in calculate("sin(pi/6)").result
    assert calculate("1001!").status == "invalid_arguments" and calculate("3 != 4").status == "invalid_arguments"


def test_a_whole_number_too_big_for_a_float_is_still_an_answer():
    """171! and 2**2000 made the finite check itself fail (OverflowError): 'an internal error'."""
    r = calculate("200!")
    assert r.status == "ok" and r.result["result"] == "7.886578673648e+374" and r.result["digits"] == 375
    assert calculate("2**2000").result["digits"] == 603 and calculate("30!").result["result"] == 265252859812191058636308480000000
    assert calculate("200! * 0.5").status == "invalid_arguments"  # a float that large can't exist: said so


def test_a_derived_table_is_a_csv_file_named_after_its_source_and_step(ws):
    """Cleaning twice once gave '(cleaned) (cleaned)'; a derived table is now a CSV file of its own, beside its source."""
    t = ws.add_table(pd.DataFrame({"a": [1]}), "sales.csv")
    cleaned = ws.add_table(pd.DataFrame({"a": [1]}), ws.derived_name(t, "cleaned"), parent=t)
    assert ws.entry(cleaned)["name"] == "sales-cleaned.csv" and ws.derived_name(cleaned, "cleaned") == "sales-cleaned.csv"
    assert ws.add_table(pd.DataFrame({"a": [2]}), ws.derived_name(cleaned, "cleaned")) != cleaned
    assert ws.label(ws.ids("table")[-1]) == "sales-cleaned (2).csv"  # the path is taken: the next one free
    assert ws.derived_name(cleaned, "predicted churn") == "sales-cleaned-predicted-churn.csv"
    assert (ws.files / "sales-cleaned.csv").read_text().splitlines() == ["a", "1"]


def test_read_code_with_line_numbers_from_a_line_on(ws):
    lines = [f"value_{i} = {i} * 2  # " + "x" * 60 for i in range(400)]
    fid = ws.add_bytes("file", "\n".join(lines).encode(), "py", name="big.py")
    first = read_file(ws, fid).result
    assert first["text"].startswith("  1| value_0 = 0 * 2") and len(first["text"]) <= HELPER_CHARS
    shown = int(first["lines"].split("-")[1].split()[0])
    assert first["next"] == f"read_file(file='big.py', start={shown + 1})"
    last = read_file(ws, fid, start=390).result
    assert last["text"].rstrip().endswith("value_399 = 399 * 2  # " + "x" * 60) and last["next"] == "end of file"
    assert read_file(ws, fid, start=999).status == "invalid_arguments"


def test_read_documents_and_notes_without_numbers(ws):
    from docx import Document
    doc = Document()
    doc.add_heading("Refund policy", 1)
    doc.add_paragraph("Opened phones can be returned within 14 days.")
    path = ws.root / "policy.docx"
    doc.save(path)
    fid = ws.add_bytes("file", path.read_bytes(), "docx", name="policy.docx")
    text = read_file(ws, fid).result["text"]
    assert "returned within 14 days" in text and "1|" not in text


def test_reading_what_isnt_text_says_which_tool_reads_it(ws):
    import io

    from PIL import Image
    t = ws.add_table(pd.DataFrame({"a": [1]}), "t")
    said = read_file(ws, t).result  # what it is, and what reads it instead
    assert said["kind"] == "table" and "table('t.csv')" in said["note"]
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "red").save(buf, format="PNG")
    img = ws.add_image(buf.getvalue(), "red.png")
    note = read_file(ws, img).result["note"]
    assert note == "An image (40x20): it has no text to read."  # a helper that sees images is shown it instead
    from app.tools.file_ops import files
    assert "view_image(image='red.png'" in files(ws, None, "read", name=img, start=1).result["note"]  # main looks
    b = ws.add_bytes("file", b"\x00\x01\x02binary", "bin", name="blob.bin")
    assert read_file(ws, b).status == "failed"


def test_listing_shows_paths_not_content(ws):
    assert ws.listing() == "Workspace: empty (no files yet)."
    t = ws.add_table(pd.DataFrame({c: [1, 2] for c in "abcdefghij"}), "sales.csv")
    ws.add_table(pd.DataFrame({"a": [1]}), ws.derived_name(t, "cleaned"), parent=t)
    ws.add_bytes("file", b"print(1)", "py", name="convert.py")
    listing = ws.listing()
    assert "sales.csv: table, 2 rows, 10 columns (a, b, c, d, e, f, g, h ...)" in listing
    assert "sales-cleaned.csv: table, 1 rows, 1 columns (a), from sales.csv" in listing
    assert "convert.py: 8 bytes" in listing and "data_" not in listing
    assert "older files not shown" in ws.listing(limit=1)


def test_sandbox_finds_files_by_name_and_old_scripts_keys(ws):
    fid = ws.add_bytes("file", b"alpha,beta\n1,2\n", "csv", name="raw.csv")
    r = run_python(ws, "print(open('raw.csv').read().splitlines()[0], open(path('raw.csv')).read().splitlines()[1])")
    assert r.result["output"].split() == ["alpha,beta", "1,2"]
    old = run_python(ws, f"print(open('{fid}').read().splitlines()[0])")  # a script saved in the first design
    assert old.result["output"].strip() == "alpha,beta"
    assert "No file 'nope.txt'. Files: raw.csv" in run_python(ws, "path('nope.txt')").result["errors"]


def test_data_files_are_read_without_line_numbers(ws):
    fid = ws.add_bytes("file", b"name,age\nAsha,31\n", "csv", name="people.csv")
    assert read_file(ws, fid).result["text"] == "name,age\nAsha,31"


def test_known_table_problems_travel_with_results(ws, tmp_path):
    from app.tools.inspect_data import inspect_data
    from app.tools.query_data import query_data
    path = tmp_path / "t.csv"
    path.write_text("id,qty\n" + "\n".join(f"{i},1" for i in range(30)) + "\nTotal,30\n", encoding="utf-8")
    data = ws.add_file(path)
    assert "data_warning" not in query_data(ws, 'SELECT sum(qty) FROM "t.csv"').result  # not inspected yet
    inspect_data(ws, data)
    warned = query_data(ws, 'SELECT sum(qty) FROM "t.csv"').result
    assert "summary rows" in warned["data_warning"] and "clean_data" in warned["data_warning"]


def test_a_table_queried_with_duckdb_in_a_script_gets_the_fix(ws):
    ws.add_table(pd.DataFrame({"a": [1, 2]}), name="t.csv")
    message = run_python(ws, "import duckdb\nduckdb.query('SELECT * FROM t').df()").result["error"]["message"]
    assert "pd.read_csv('its name') or table('its name')" in message
    assert "In a script" not in run_python(ws, "1 / 0").result["error"]["message"]


def test_a_script_changes_the_assistants_files_with_versions_and_never_the_users(ws, tmp_path):
    fid = ws.add_bytes("file", b"ORIGINAL", "py", name="average.py")
    memory = agent_md.ensure(ws)
    upload = tmp_path / "t.csv"
    upload.write_text("a\n1\n")
    table = ws.add_file(upload)  # the user's own
    changed = run_python(ws, "print(open('average.py').read())\nopen('average.py', 'w').write('CHANGED')")
    assert changed.result["output"].strip() == "ORIGINAL" and ws.path(fid).read_bytes() == b"CHANGED"
    assert changed.result["files"] == [{"path": "average.py", "kind": "file", "updated": True}] and len(ws.versions(fid)) == 1
    refused = run_python(ws, "open('t.csv', 'w').write('a\\n2\\n')")
    assert refused.status == "invalid_arguments" and "the user's own file" in refused.result["errors"]
    assert ws.table(table)["a"].tolist() == [1]
    kept = run_python(ws, "open('agent.md', 'a').write('- run by a script')")  # agent.md is the project's memory
    assert kept.status == "invalid_arguments" and "script" not in agent_md.text(ws) and memory in ws.ids()


def test_env_file_fills_the_environment_without_overriding(tmp_path, monkeypatch):
    from app.config import load_env_file
    env = tmp_path / ".env"
    env.write_text('# secrets\nCC_TEST_KEY="sk_abc"\nexport CC_TEST_OTHER = plain\nCC_TEST_SET=from-file\nnot a line\n')
    monkeypatch.delenv("CC_TEST_KEY", raising=False)
    monkeypatch.delenv("CC_TEST_OTHER", raising=False)
    monkeypatch.setenv("CC_TEST_SET", "from-environment")
    load_env_file(env)
    import os
    assert os.environ["CC_TEST_KEY"] == "sk_abc" and os.environ["CC_TEST_OTHER"] == "plain"
    assert os.environ["CC_TEST_SET"] == "from-environment"  # a variable already set wins
    load_env_file(tmp_path / "missing.env")  # no file: nothing happens
    monkeypatch.delenv("CC_TEST_KEY")
    monkeypatch.delenv("CC_TEST_OTHER")
