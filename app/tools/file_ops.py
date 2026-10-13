"""files(): the one tool for the workspace's files - for main and for the helpers (who get find, read, write and
edit): list, find, read, write, edit, rename, delete, download.

    write   a new file - or, for a file that is there, its whole new text (the version before is kept)
    edit    one passage changed: old_text exactly as it is in the file -> new_text (a file shown with line numbers
            may be quoted with them: they are the reader's, not the file's)
    rename  a file moved to another path

A file is named by its path ('report.md', 'charts/sales.png'), in any case, or a copied listing line - the
workspace's one resolver (Workspace.resolve); a name that isn't there gets the closest paths back, so an invented
name is corrected in one step. read gives a file's report first (what a helper did with it) when it has one; with
start, its content - the one reader, read_file.read.

The user decides (approval.py): delete and download always ask with a button, and so does changing a file the user
brought (an upload or a download) - it is theirs, even though its version before is kept. Files the assistant made
change freely; renaming loses nothing and never asks. A .docx (headings, paragraphs, lists, tables, bold and italic)
and a .pptx (a slide per # heading: its bullets, a picture, speaker notes) are written from Markdown by code; a line
that is only ![caption](charts/sales.png) puts that picture in - a generated image, or a chart (every chart is also
kept as a picture, chart_image.py).
"""

from __future__ import annotations

import difflib
import io
import re
import tempfile
from pathlib import Path, PurePath

from . import approval, read_file
from .core import ToolError, ToolResult, ok, tool
from .filetypes import BINARY, FROM_MARKDOWN, is_code, suffix
from .search_index import SearchIndex
from .workspace import Workspace, clean_path

PAGE = 15  # items per list page
KINDS = ("table", "image", "file")  # everything but notes
MAX_WRITE = 1_000_000  # characters
MODES = ("list", "find", "read", "write", "edit", "rename", "delete", "download")
HELPER_MODES = ("find", "read", "write", "edit")  # a helper's files: what it reads and makes; asking is main's
# Formats that are not text (filetypes.BINARY): text written under such a name would be a broken file (a .docx and
# a .pptx are the exceptions - FROM_MARKDOWN, built from Markdown). A script makes them (analyze_or_code).


@tool
def files(ws: Workspace, index: SearchIndex | None, mode: str, *, name: str | None = None, page: int = 1,
          start: int | None = None, content: str | None = None, old_text: str | None = None,
          new_text: str | None = None, new_name: str | None = None, link: str | None = None,
          helper: bool = False) -> ToolResult:
    mode = str(mode).strip().lower()
    modes = HELPER_MODES if helper else MODES
    if mode not in modes:
        raise ToolError(f"mode must be one of {', '.join(modes)}.", field="mode")
    if mode == "list":
        return _list(ws, page)
    if mode == "download":
        return _download(_need(link, "link", "download needs link: a GitHub file's link, or a data file's link"))
    name = _need(name, "name", f"{mode} needs name: a file's name, e.g. 'report.md'")
    if mode == "find":
        return _find(ws, name)
    if mode == "read":
        return _read(ws, _resolve(ws, name), start, helper)
    if mode == "write":
        return _write(ws, index, name, _need(content, "content", "write needs content: the file's whole text"))
    if mode == "edit":
        if new_text is None:
            raise ToolError("edit needs new_text: what replaces old_text ('' removes it).", field="new_text")
        return _edit(ws, index, _resolve(ws, name), _need(old_text, "old_text", "edit needs old_text: the passage "
                                                                               "to change, exactly as in the file"),
                     str(new_text))
    if mode == "rename":
        return _rename(ws, index, _resolve(ws, name), _need(new_name, "new_name", "rename needs new_name"))
    return approval.delete_request(ws, [_resolve(ws, name)])


# ---------------------------------------------------------------- list, find, read

def _list(ws: Workspace, page: int) -> ToolResult:
    lines = list(reversed(ws.lines().values()))  # newest first
    if not lines:
        return ok({"items": [], "note": "The workspace is empty: nothing uploaded, downloaded or made yet."})
    pages = (len(lines) + PAGE - 1) // PAGE
    if not 1 <= page <= pages:
        raise ToolError(f"There are {pages} page(s); page must be 1 to {pages}.", field="page")
    shown = lines[(page - 1) * PAGE:page * PAGE]
    more = f"more: files(mode='list', page={page + 1})" if page < pages else "end of list"
    return ok({"total": len(lines), "page": f"{page} of {pages}", "items": shown, "next": more})


def _find(ws: Workspace, name: str) -> ToolResult:
    item = _exact(ws, name)
    if item:
        return ok({"found": True, "item": ws.lines()[item]})
    try:
        _resolve(ws, name)
    except ToolError as missing:  # the closest items, as 'id (name)'
        return ok({"found": False, "note": missing.message})


def _read(ws: Workspace, item: str, start: int | None, helper: bool = False) -> ToolResult:
    """The one reader (read_file.read). main: the report alone first, then the text from start; a helper: the report
    beside the first lines, and more of them at once (it reads a long document in fewer steps)."""
    if helper:
        return read_file.read(ws, item, start, read_file.HELPER_CHARS, tool="files")
    return read_file.read(ws, item, start, read_file.MAIN_CHARS, tool="files", report_alone=True)


# ---------------------------------------------------------------- write

def _write(ws: Workspace, index: SearchIndex | None, name: str, content: str) -> ToolResult:
    """A new file - or, for one that is there, its whole new text (the user's own file asks first)."""
    _fits(content)
    existing = _exact(ws, name)  # a path, as given - before a default suffix is added
    if existing is None:
        name = _file_name(name)
        existing = _exact(ws, name)
    if existing is None:
        _writable(name)
        _pictures_there(ws, name, content)
        return _store_new(ws, index, name, content)
    return _change(ws, index, existing, content)


def _change(ws: Workspace, index: SearchIndex | None, item: str, content: str) -> ToolResult:
    """A file's new text: the assistant's own is changed (its version before kept), the user's asks first."""
    e = ws.entry(item)
    if e["kind"] == "image" or (e["kind"] == "table" and suffix(e["name"]) not in (".csv", ".tsv")):
        raise ToolError(f"{e['name']} is {'an image' if e['kind'] == 'image' else 'a spreadsheet'}, which isn't "
                        "written as text: write a new file, or a script makes it.", field="name")
    _writable(e["name"])  # before the user is asked to approve it: a PDF written as text would be broken
    _pictures_there(ws, e["name"], content)
    if made_by_assistant(e):
        return replace_file(ws, index, item, content)
    return approval.replace_request(ws, item, content)


GUTTER = re.compile(r"^\s*(?:>>)?\s*\d+\| ?", re.M)  # the line numbers files(read) and a run's 'around' show


def _edit(ws: Workspace, index: SearchIndex | None, item: str, old: str, new: str) -> ToolResult:
    """One passage changed: ``old`` exactly as it is in the file (found once), replaced by ``new``."""
    e = ws.entry(item)
    if e["name"].lower().endswith(FROM_MARKDOWN):
        raise ToolError(f"A {PurePath(e['name']).suffix} is written whole from Markdown: give its full text with "
                        "mode='write'.", field="name")
    if e["kind"] == "image" or suffix(e["name"]) in BINARY:
        raise ToolError(f"{e['name']} isn't text, so it can't be edited: a script changes it.", field="name")
    text = ws.path(item).read_text("utf-8", errors="replace")
    if text.count(old) == 0 and all(GUTTER.match(ln) for ln in old.splitlines() if ln.strip()):
        old, new = GUTTER.sub("", old), GUTTER.sub("", new)  # quoted as read, with the reader's line numbers
    found = text.count(old)
    if found == 0:
        close = difflib.get_close_matches(old.splitlines()[0] if old.strip() else old, text.splitlines(), n=3, cutoff=0.5)
        near = (" The closest lines: " + " | ".join(c.strip() for c in close) + ".") if close else ""
        raise ToolError(f"old_text is not in {e['name']}: copy the passage exactly as it is in the file (read it with "
                        f"files(mode='read')), spaces and line breaks included.{near}", field="old_text")
    if found > 1:
        raise ToolError(f"old_text is in {e['name']} {found} times: give more of the lines around it, so it is found "
                        "once.", field="old_text")
    return _change(ws, index, item, text.replace(old, new, 1))


def _rename(ws: Workspace, index: SearchIndex | None, item: str, new_name: str) -> ToolResult:
    from .files import index_report, reindex
    e = ws.entry(item)
    old = e["name"]
    new_name = " ".join(str(new_name).split())
    if not PurePath(new_name).suffix and PurePath(old).suffix:
        new_name += PurePath(old).suffix  # 'report' for 'notes.md' keeps .md
    _file_name(new_name)
    if new_name == old:  # measured: 'tts.rs -> tts.rs' was reported as a rename
        raise ToolError(f"{item} is already called {old}: nothing to rename.", field="new_name")
    if PurePath(new_name).suffix.lower() != PurePath(old).suffix.lower() and e["kind"] == "file":
        _writable(new_name)  # notes.md renamed notes.pdf would still be text: a broken PDF
    taken = _exact(ws, new_name)
    if taken and taken != item:
        raise ToolError(f"{new_name} is already a file here: pick another name.", field="new_name")
    ws.rename(item, new_name)
    if index is not None and "passages" in e.get("meta", {}):
        reindex(ws, index, item)  # its passages carry its name
    index_report(ws, index, item)
    return ok({"file": item, "renamed": f"{old} -> {new_name}"})


def _fits(content: str) -> None:
    if len(content) > MAX_WRITE:
        raise ToolError(f"content is longer than {MAX_WRITE:,} characters.", field="content")


def _store_new(ws: Workspace, index: SearchIndex | None, name: str, content: str) -> ToolResult:
    from .files import store_file
    data = _bytes(name, content, ws)
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / PurePath(name).name  # kept under its whole path, folder and all
        path.write_bytes(data)
        kept = store_file(ws, index, path, by="assistant", name=name)
    item = kept.get("id") or kept.get("file")
    return ok({"file": item, "name": kept["name"], "kind": kept["kind"], "bytes": len(data),
               "note": "Saved; the user sees it in the chat.", **written_code(name, content)}, _card(ws, item, "written"))


def replace_file(ws: Workspace, index: SearchIndex | None, item: str, content: str) -> ToolResult:
    """Replace a stored file's content (and its search passages, for a document)."""
    from .files import reindex
    name = ws.entry(item)["name"]
    data = _bytes(name, content, ws)
    ws.replace_bytes(item, data)
    if index is not None and "passages" in ws.entry(item).get("meta", {}):
        reindex(ws, index, item)
    return ok({"file": item, "name": name, "bytes": len(data), "note": "Changed (its version before is kept); the user "
               "sees it in the chat.", **written_code(name, content)}, _card(ws, item, "updated"))


def written_code(name: str, content: str) -> dict:
    """For a program written or changed as text: it was not run - and, for Python, whether it still parses, with the
    line where it breaks. Measured on the 26B: asked to fix a crashing stats.py and say what it prints, main
    changed one line with files(update), left a second 'def mean(values):' in it, and answered 'It now prints:
    mean reading: 15.1' - a run that never happened, of a file that still didn't run."""
    if not is_code(name):
        return {}
    out = {"not_run": f"{name} was written, not run: don't say what it prints before a run (run runs it in a helper; "
                      "main gives it to analyze_or_code)."}
    if suffix(name) == ".py":
        try:
            compile(content, name, "exec")
        except SyntaxError as exc:
            out["broken"] = f"{name} does not parse: {exc.msg} at line {exc.lineno}. Fix it before saying it works."
    return out


def made_by_assistant(entry: dict) -> bool:
    return entry.get("by") == "assistant"


def _card(ws: Workspace, item: str, action: str) -> dict:
    e = ws.entry(item)
    return {"type": "file", "id": item, "name": e["name"], "kind": e["kind"], "bytes": e.get("bytes"),
            "action": action}


def _file_name(name: str) -> str:
    """A path for a new file, from the model: refused when it isn't one (workspace.clean_path); Markdown when it
    has no extension."""
    name = clean_path(name, strict=True)
    return name if PurePath(name).suffix else name + ".md"


def _writable(name: str) -> None:
    """Text can be written under this name: a text format, or .docx (built from Markdown)."""
    suffix = PurePath(name).suffix.lower()
    if suffix in BINARY:
        raise ToolError(f"A {suffix} file isn't text, so it can't be written or changed here - the text would be a "
                        f"broken {suffix} file. Write .md, .txt, .csv, code, or .docx and .pptx (from Markdown); a "
                        f"{suffix} file is made by a script (analyze_or_code).", field="name")


def _bytes(name: str, content: str, ws: Workspace | None = None) -> bytes:
    _writable(name)
    picture = (lambda ref: _picture(ws, ref)) if ws is not None else (lambda ref: None)
    if name.lower().endswith(".docx"):
        return markdown_docx(content, picture)
    if name.lower().endswith(".pptx"):
        return markdown_pptx(content, picture)
    return content.encode("utf-8")


def _picture(ws: Workspace, ref: str) -> Path | None:
    """The stored picture a Markdown image names - by id ('image_7') or name; None when there is none."""
    item = ws.find(ref, ("image",))
    return ws.path(item) if item else None


PICTURE = re.compile(r"^\s*!\[([^\]]*)\]\(\s*<?([^)>]+?)>?\s*\)\s*$")


def _missing_picture(ref: str) -> str:
    """What a document shows for a picture that isn't there - only for the user's own edits (the UI): the
    model's write and update are refused first (_pictures_there), so it never saves a broken document."""
    return f"[picture {ref} not found]"


def _pictures_there(ws: Workspace, name: str, content: str) -> None:
    """A .docx or .pptx whose Markdown names a picture that isn't in the workspace is refused, with the
    pictures that are - instead of being saved with '[picture ... not found]' in it while the result says
    'Saved'."""
    if not name.lower().endswith(FROM_MARKDOWN):
        return
    refs = [m[2].strip() for m in map(PICTURE.match, content.splitlines()) if m]
    missing = list(dict.fromkeys(r for r in refs if ws.find(r, ("image",)) is None))
    if not missing:
        return
    images = ws.ids("image")[-10:]
    have = (", ".join(ws.label(i) for i in images) if images
            else "none - a chart from analyze_or_code or a picture from generate_image is kept as one")
    raise ToolError(f"No picture {', '.join(missing)} in the workspace, so nothing was written. Pictures: {have}. "
                    "Name one of them in ![caption](its name), or leave the picture out.", field="content")


# ---------------------------------------------------------------- download

def _download(link: str) -> ToolResult:
    """One GitHub file of any kind, or a data table from any public site - both ask the user first."""
    from .web import _on_github, _raw_link
    if _on_github(_raw_link(link.strip())):
        return approval.download_request(link)
    return approval.import_request(link)


# ---------------------------------------------------------------- names

def _need(value, field: str, message: str) -> str:
    if value is None or not str(value).strip():
        raise ToolError(message + ".", field=field)
    return str(value) if field == "content" else str(value).strip()


def _exact(ws: Workspace, name: str) -> str | None:
    """The item this names (id, name or a copied listing line), or None."""
    return ws.find(name, KINDS)


def _resolve(ws: Workspace, name: str) -> str:
    return ws.resolve(name, KINDS, field="name")


# ---------------------------------------------------------------- .docx from Markdown

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def markdown_docx(text: str, picture=lambda ref: None) -> bytes:
    """A Word document from Markdown: # headings, paragraphs, - and 1. lists, | tables |, **bold**,
    *italic*, `code`, and a picture on a line of its own: ![caption](image_7) (``picture`` finds its file)."""
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for r, cells in enumerate(rows):
                    for c, cell in enumerate(cells):
                        _runs(table.cell(r, c).paragraphs[0], cell, bold=r == 0)
            continue
        heading = re.match(r"(#{1,4})\s+(.*)", line)
        bullet = re.match(r"\s*[-*+]\s+(.*)", line)
        number = re.match(r"\s*\d+[.)]\s+(.*)", line)
        image = PICTURE.match(line)
        if image:
            path = picture(image[2].strip())
            if path is None:
                _runs(doc.add_paragraph(), _missing_picture(image[2].strip()))
            else:
                doc.add_picture(str(path), width=Inches(6))
                if image[1].strip():
                    doc.add_paragraph(image[1].strip()).runs[0].italic = True
        elif heading:
            doc.add_heading(heading[2].strip(), level=len(heading[1]))
        elif bullet:
            _runs(doc.add_paragraph(style="List Bullet"), bullet[1])
        elif number:
            _runs(doc.add_paragraph(style="List Number"), number[1])
        elif line.strip():
            _runs(doc.add_paragraph(), line.strip())
        i += 1
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def markdown_pptx(text: str, picture=lambda ref: None) -> bytes:
    """A slide deck (16:9) from Markdown: each # (or ##) heading starts a slide with that title; its lines are
    its bullets (- or 1. or plain; indented ones a level down; **bold** kept), a line that is only
    ![caption](image_7) its picture (beside the bullets, or alone, large), a 'Notes:' line its speaker notes.
    A first slide with no more than one line under its heading is a title slide (that line its subtitle)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slides, current = [], None
    for raw in text.splitlines():
        line = raw.rstrip()
        heading = re.match(r"#{1,2}\s+(.*)", line)
        if heading:
            current = {"title": heading[1].strip(), "lines": [], "pictures": [], "notes": []}
            slides.append(current)
            continue
        if not line.strip() or re.fullmatch(r"\s*(?:-{3,}|\*{3,})\s*", line):
            continue
        if current is None:  # text before any heading: a slide without a title
            current = {"title": "", "lines": [], "pictures": [], "notes": []}
            slides.append(current)
        image = PICTURE.match(line)
        if image:
            current["pictures"].append((image[1].strip(), image[2].strip()))
        elif re.match(r"\s*notes?:", line, re.I):
            current["notes"].append(line.split(":", 1)[1].strip())
        else:
            level = min(2, (len(line) - len(line.lstrip())) // 2)
            current["lines"].append((level, re.sub(r"^\s*(?:[-*+]|\d+[.)]|#{3,})\s*", "", line)))
    if not slides:
        raise ToolError("The slides need content: start each slide with a # heading.", field="content")
    width, height = deck.slide_width, deck.slide_height
    for n, s in enumerate(slides):
        if n == 0 and len(s["lines"]) <= 1 and not s["pictures"]:
            slide = deck.slides.add_slide(deck.slide_layouts[0])  # the title slide
            slide.shapes.title.text = s["title"]
            if s["lines"]:
                slide.placeholders[1].text = s["lines"][0][1].replace("**", "")
            else:
                slide.placeholders[1].element.getparent().remove(slide.placeholders[1].element)
        else:
            slide = deck.slides.add_slide(deck.slide_layouts[5])  # a title only: the rest placed here
            slide.shapes.title.text = s["title"]
            top, bottom, left = Inches(1.7), height - Inches(0.5), Inches(0.6)
            text_width = width - 2 * left if not s["pictures"] else (width - 2 * left) * 0.45 if s["lines"] else 0
            if s["lines"]:
                frame = slide.shapes.add_textbox(left, top, text_width, bottom - top).text_frame
                frame.word_wrap = True
                for k, (level, words) in enumerate(s["lines"]):
                    p = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
                    p.level = level
                    _slide_runs(p, ("• " if level == 0 else "– ") + words)
                    for run in p.runs:
                        run.font.size = Pt(22 if level == 0 else 18)
            for k, (caption, ref) in enumerate(s["pictures"][:1]):  # one picture a slide: it should be readable
                path = picture(ref)
                box_left = left + text_width + (Inches(0.3) if s["lines"] else 0)
                box_width, box_height = width - box_left - left, bottom - top - (Inches(0.5) if caption else 0)
                if path is None:
                    slide.shapes.add_textbox(box_left, top, box_width, Inches(1)).text_frame.text = _missing_picture(ref)
                    continue
                shape = slide.shapes.add_picture(str(path), box_left, top)
                scale = min(box_width / shape.width, box_height / shape.height)
                shape.width, shape.height = int(shape.width * scale), int(shape.height * scale)
                shape.left = int(box_left + (box_width - shape.width) / 2)
                if caption:
                    cap = slide.shapes.add_textbox(box_left, top + shape.height + Inches(0.1), box_width, Inches(0.4))
                    cap.text_frame.text = caption
                    cap.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
                    cap.text_frame.paragraphs[0].runs[0].font.italic = True
        if s["notes"]:
            slide.notes_slide.notes_text_frame.text = "\n".join(s["notes"])
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def _runs(paragraph, text: str, bold: bool = False) -> None:
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            paragraph.add_run(piece[1:-1]).font.name = "Consolas"
        else:
            paragraph.add_run(piece).bold = bold or None


def _slide_runs(paragraph, text: str) -> None:
    """_runs for a slide's paragraph (python-pptx: a run gets its text after it is added; bold is on its font)."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        run = paragraph.add_run()
        if piece.startswith("**") and piece.endswith("**"):
            run.text, run.font.bold = piece[2:-2], True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run.text, run.font.italic = piece[1:-1], True
        elif piece.startswith("`") and piece.endswith("`"):
            run.text, run.font.name = piece[1:-1], "Consolas"
        else:
            run.text = piece
