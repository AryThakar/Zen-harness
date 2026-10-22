# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0
"""Text out of documents: PDF, Word (.docx), PowerPoint (.pptx), plain text and Markdown - the one
reader for uploads, files(read) and web links to documents.

Each document becomes a list of blocks - a paragraph (or table row) with the PDF
page or slide it is on and the heading it sits under - so search results can say
where a passage came from. lines() writes a document out as text for reading.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import PurePath

from app.tools.filetypes import DOCUMENT as KINDS  # noqa: E402 - the one definition of a document


class DocumentError(ValueError):
    """The file can't be used; the message says why in plain words."""


@dataclass
class Block:
    text: str
    page: int | None = None
    section: str | None = None


@dataclass
class Extracted:
    kind: str
    title: str
    blocks: list[Block]
    pages: int | None = None

    @property
    def words(self) -> int:
        return sum(len(b.text.split()) for b in self.blocks)


def kind_of(filename: str) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix not in KINDS:
        raise DocumentError(f"'{suffix or filename}' files aren't supported; use PDF, Word (.docx), PowerPoint "
                            "(.pptx), .txt or .md.")
    return KINDS[suffix]


def extract(filename: str, data: bytes) -> Extracted:
    kind = kind_of(filename)  # any size: it is read into passages (uploads have their limit: BIZ_MAX_UPLOAD_MB)
    if not data:
        raise DocumentError("The file is empty.")
    stem = PurePath(filename).stem
    if kind == "pdf":
        result = _pdf(data, stem)
    elif kind == "docx":
        result = _docx(data, stem)
    elif kind == "pptx":
        result = _pptx(data, stem)
    else:
        result = _text(_decode(data), stem, markdown=kind == "md")
    result.blocks = [b for b in result.blocks if b.text]
    if not result.blocks:
        hint = " It may be a scanned PDF (images of pages); text recognition isn't supported." if kind == "pdf" else ""
        raise DocumentError(f"No text found in the file.{hint}")
    return result


# --- formats ---------------------------------------------------------------------

def _pdf(data: bytes, stem: str) -> Extracted:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentError("The PDF is password-protected.")
        pages = [page.extract_text() or "" for page in reader.pages]
        title = (reader.metadata.title if reader.metadata else None) or stem
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise DocumentError(f"The PDF couldn't be read ({exc}).") from exc
    blocks = [Block(p, page=n) for n, text in enumerate(pages, start=1) for p in _paragraphs(text)]
    return Extracted("pdf", _clean(str(title)), blocks, pages=len(pages))


def _docx(data: bytes, stem: str) -> Extracted:
    import zipfile

    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, ValueError, KeyError) as exc:
        raise DocumentError("The Word file couldn't be read; save it as .docx and try again.") from exc
    blocks: list[Block] = []
    headings: list[Block] = []  # used only if the document has nothing else
    section = None
    body = document.element.body
    paragraphs = {p._element: p for p in document.paragraphs}
    tables = {t._element: t for t in document.tables}
    for child in body.iterchildren():  # paragraphs and tables in document order
        if child in paragraphs:
            p = paragraphs[child]
            text = _clean(p.text)
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") or style == "title":
                section = text  # kept as the passages' section, not repeated in their text
                headings.append(Block(text, section=text))
                continue
            blocks.append(Block(text, section=section))
        elif child in tables:
            for row in tables[child].rows:
                cells = []
                for cell in row.cells:  # merged cells repeat; keep each once
                    value = _clean(cell.text)
                    if value and (not cells or cells[-1] != value):
                        cells.append(value)
                if cells:
                    blocks.append(Block(" | ".join(cells), section=section))
    title = _clean(document.core_properties.title or "") or stem
    return Extracted("docx", title, blocks or headings)


def _pptx(data: bytes, stem: str) -> Extracted:
    """Each slide's text - its title as the passages' section, then its text boxes and table rows in order;
    a slide's number is its page."""
    import zipfile

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        deck = Presentation(io.BytesIO(data))
    except (zipfile.BadZipFile, ValueError, KeyError) as exc:
        raise DocumentError("The PowerPoint file couldn't be read; save it as .pptx and try again.") from exc

    def shapes(group):
        for shape in group:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from shapes(shape.shapes)
            else:
                yield shape

    blocks: list[Block] = []
    for n, slide in enumerate(deck.slides, start=1):
        title_shape = slide.shapes.title
        section = (_clean(title_shape.text_frame.text) if title_shape is not None else "") or f"Slide {n}"
        for shape in shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = _clean("".join(run.text for run in para.runs))
                    if text:
                        blocks.append(Block(text, page=n, section=section))
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [_clean(cell.text) for cell in row.cells if _clean(cell.text)]
                    if cells:
                        blocks.append(Block(" | ".join(cells), page=n, section=section))
        if not any(b.page == n for b in blocks) and section != f"Slide {n}":
            blocks.append(Block(section, page=n, section=section))  # a slide with only its title
    title = _clean(deck.core_properties.title or "") or stem
    return Extracted("pptx", title, blocks, pages=len(deck.slides))


def lines(doc: Extracted) -> list[str]:
    """A document written out for reading: a [page N] (or [slide N]) line where the page changes, a heading
    line where the section changes, then each block's text."""
    out, page, section = [], None, None
    marker = "slide" if doc.kind == "pptx" else "page"
    for b in doc.blocks:
        if b.page is not None and b.page != page:
            page = b.page
            out.append(f"[{marker} {page}]")
        if b.section and b.section != section and b.text != b.section:
            section = b.section
            out.append(f"## {section}")
        out.append(b.text)
    return out


def _text(text: str, stem: str, markdown: bool) -> Extracted:
    blocks, headings, section, title = [], [], None, None
    for para in _paragraphs(text):
        heading = re.match(r"^#{1,6}\s+(.*)", para) if markdown else None
        if heading:  # kept as the passages' section, not repeated in their text
            section = heading.group(1).strip()
            title = title or section
            headings.append(Block(section, section=section))
            continue
        blocks.append(Block(para, section=section))
    return Extracted("md" if markdown else "txt", title or stem, blocks or headings)


# --- helpers ---------------------------------------------------------------------

def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _paragraphs(text: str) -> list[str]:
    """Split on blank lines; inside a paragraph, join the lines a PDF wraps."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    parts = re.split(r"\n\s*\n", text)
    out = []
    for part in parts:
        lines = [ln.strip() for ln in part.split("\n") if ln.strip()]
        if not lines:
            continue
        # Keep list items and headings on their own lines; join wrapped prose.
        merged: list[str] = []
        for line in lines:
            starts_item = re.match(r"^([-*•]|\d+[.)]|#{1,6}\s)", line)
            if merged and not starts_item and not re.match(r"^#{1,6}\s", merged[-1]):
                merged[-1] = merged[-1][:-1] + line if merged[-1].endswith("-") else f"{merged[-1]} {line}"
            else:
                merged.append(line)
        out.extend(_clean(m) for m in merged)
    return [p for p in out if p]


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()
